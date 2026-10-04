import os
from io import BytesIO
from openpyxl import Workbook, load_workbook

os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.setdefault("ADMIN_USER", "admin")
os.environ.setdefault("ADMIN_PASSWORD", "password")
os.environ.setdefault("GOOGLE_DRIVE_ENABLED", "false")

from app import create_empty_planning_blocks, latest_planning_by_subject_group

planning = [
    {"subject": "Matemáticas", "group": "CLEI II", "theme": "Fracciones", "week": "Semana 2", "week_number": 2, "date_label": "2 al 6 de febrero", "status": "DICTADA"},
    {"subject": "Matemáticas", "group": "CLEI II", "theme": "Ecuaciones", "week": "Semana 4", "week_number": 4, "date_label": "16 al 20 de febrero", "status": "DICTADA", "context": "Alta"},
    {"subject": "Matemáticas", "group": "CLEI II", "theme": "Tema futuro pendiente", "week": "Semana 99", "week_number": 99, "date_label": "futuro", "status": "PENDIENTE", "context": "Alta"},
    {"subject": "Matemáticas", "group": "CLEI III", "theme": "Álgebra", "week": "Semana 3", "week_number": 3, "date_label": "9 al 13 de febrero", "status": "DICTADA"},
    {"subject": "Biología", "group": "CLEI II", "theme": "La célula", "week": "Semana 5", "week_number": 5, "date_label": "23 al 27 de febrero", "status": "DICTADA"},
    {"subject": "Biología", "group": "CLEI III", "theme": "Sin planeación registrada", "week": "Semana 8", "week_number": 8, "date_label": "marzo", "status": "N/A"},
    {"subject": "Matemáticas", "group": "CLEI I", "theme": "Tema que no debe aparecer", "week": "Semana 99", "week_number": 99, "date_label": "", "status": "DICTADA"},
    {"subject": "Matemáticas", "group": "MULTIGRADO / CLEI II", "theme": "Tema de Multigrado que no debe aparecer", "week": "Semana 100", "week_number": 100, "date_label": "", "status": "DICTADA", "context": "Multigrado", "source": "Multigrado"},
]

latest = latest_planning_by_subject_group(planning)
assert latest[("Matemáticas", "CLEI II")]["theme"] == "Tema futuro pendiente"
assert latest[("Matemáticas", "CLEI II")]["week_number"] == 99
assert latest[("Matemáticas", "CLEI III")]["theme"] == "Álgebra"
assert latest[("Biología", "CLEI II")]["theme"] == "La célula"
assert ("Biología", "CLEI III") not in latest
assert not any(clei == "CLEI I" for _, clei in latest)

workbook = Workbook()
worksheet = workbook.active
worksheet.title = "Matemáticas"
worksheet.cell(1, 2).value = "Semana 4"
worksheet.cell(1, 3).value = "16 al 20 de febrero de 2026"
for offset, clei in enumerate(("CLEI I", "CLEI II", "CLEI III", "CLEI IV", "CLEI V", "CLEI VI")):
    worksheet.cell(1 + offset, 4).value = clei
    worksheet.cell(1 + offset, 5).value = "Tema anterior"
    worksheet.cell(1 + offset, 6).value = "Página 1"
    worksheet.cell(1 + offset, 7).value = "Objetivo anterior"
    worksheet.cell(1 + offset, 8).value = "Actividad anterior"
    worksheet.cell(1 + offset, 9).value = "DICTADA"
buffer = BytesIO()
workbook.save(buffer)
updated, created, _ = create_empty_planning_blocks(buffer, ["Matemáticas"], selected_cleis={"Matemáticas": []})
check = load_workbook(updated, data_only=True)["Matemáticas"]
assert [check.cell(7, column).value for column in range(5, 10)] == ["N/A"] * 5

print("OK: última clase registrada aislada por materia/CLEI y CLEI 1 queda en N/A.")
