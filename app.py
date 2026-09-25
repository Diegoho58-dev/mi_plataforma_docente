import io
import json
import os
import re
import time
import unicodedata
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from functools import wraps

from flask import Flask, make_response, redirect, render_template, request, session, url_for
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from gemini_planning import GeminiPlanningError, generate_planning_proposal
from external_subjects import build_external_matrix, consolidate_records, mark_week_mismatches, parse_values, summarize

app = Flask(__name__)

SECRET_KEY = os.environ.get("SECRET_KEY")
ADMIN_USER = os.environ.get("ADMIN_USER")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD")

if not SECRET_KEY:
    raise RuntimeError("Falta la variable de entorno SECRET_KEY.")
if not ADMIN_USER or not ADMIN_PASSWORD:
    raise RuntimeError("Faltan las variables ADMIN_USER o ADMIN_PASSWORD.")

app.config["SECRET_KEY"] = SECRET_KEY

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive", "https://www.googleapis.com/auth/spreadsheets"]
PLANNING_SHEETS = {"tecnico laboral", "comunidad terapeutica", "maxima", "multigrado"}
STUDENT_SHEETS = {"clei 2", "clei 3a", "clei3b", "clei 4", "clei 5-6", "mult. asistencia"}
COMMUNITY_SHEET_MARKER = "comunidad terapeutica"
COMMUNITY_NOTES_SHEET_MARKER = "notas com ter"
CONTEXT_OPTIONS = ["Alta", "Multigrado", "Técnico Laboral", "Comunidad Terapéutica"]
CYCLE_START = date(2026, 7, 6)
DRIVE_ENABLED = os.environ.get("GOOGLE_DRIVE_ENABLED", "false").strip().lower() == "true"
PLANNING_ONLY_FRIDAY = os.environ.get("PLANNING_ONLY_FRIDAY", "false").strip().lower() == "true"
DEFAULT_PLANNING_DRIVE_FILE_ID = "1qNzaB4pFeNUUQPRJ48Ay-afEuwvxbPCu"
EXTERNAL_HIGH_CLEI_FILE_ID = os.environ.get("EXTERNAL_HIGH_CLEI_FILE_ID", "1i8fwoPaqB7vzw07Xt2K4oDji4QH3cMuTj9O1xxG8Cjw")
EXTERNAL_MULTIGRADE_FILE_ID = os.environ.get("EXTERNAL_MULTIGRADE_FILE_ID", "1i1T_NlN_j7DAt0GNnfSVeR4dGRol_36ysDdwLFU_3KQ")
EXTERNAL_SUBJECTS_ENABLED = os.environ.get("EXTERNAL_SUBJECTS_ENABLED", "true").strip().lower() == "true"
CURRICULUM_CACHE = {"modified_time": None, "topics": None, "pages": None, "loaded_at": 0}
CURRICULUM_CACHE_SECONDS = int(os.environ.get("CURRICULUM_CACHE_SECONDS", "300"))
EXTERNAL_SUBJECTS_CACHE = {"records": None, "metadata": [], "loaded_at": 0}
EXTERNAL_SUBJECTS_CACHE_SECONDS = int(os.environ.get("EXTERNAL_SUBJECTS_CACHE_SECONDS", "300"))


def login_required(view_function):
    @wraps(view_function)
    def protected_view(*args, **kwargs):
        if "user" not in session:
            return redirect(url_for("login"))
        return view_function(*args, **kwargs)
    return protected_view


def clean_text(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split()).strip()


def format_date(value):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        days = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
        months = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
        return f"{days[value.weekday()]} {value.day} de {months[value.month - 1]} de {value.year}"
    return clean_text(value)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value).lower()
    match = re.search(r"(\d{1,2})\s+de\s+([a-záéíóú]+)\s+de\s+(\d{4})", text)
    if not match:
        return None
    months = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
    month = months.get(match.group(2))
    if not month:
        return None
    try:
        return date(int(match.group(3)), month, int(match.group(1)))
    except ValueError:
        return None


def cycle_for_date(value):
    class_date = parse_date(value)
    if not class_date or class_date < CYCLE_START:
        return None
    week_number = ((class_date - CYCLE_START).days // 7) + 1
    return {"cycle": ((week_number - 1) // 3) + 1, "week": ((week_number - 1) % 3) + 1}


def current_cycle():
    info = cycle_for_date(date.today())
    return info or {"cycle": 1, "week": 1}


def is_no_class(theme, observations):
    text = f"{theme} {observations}".lower()
    markers = ("no hubo clase", "no hay clase", "sin clase", "no se realizó", "no se realizo", "no se dictó", "no se dicto", "suspendida")
    return any(marker in text for marker in markers)


def classify_context(sheet_name, group):
    text = normalize_header(f"{sheet_name} {group}")
    if "tecnico laboral" in text:
        return "Técnico Laboral"
    if "comunidad terapeutica" in text:
        return "Comunidad Terapéutica"
    return "Multigrado" if "multigrado" in text else "Alta"


def get_drive_service():
    file_id = os.environ.get("GOOGLE_DRIVE_FILE_ID")
    service_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not file_id or not service_json:
        raise RuntimeError("Faltan GOOGLE_DRIVE_FILE_ID o GOOGLE_SERVICE_ACCOUNT_JSON en Render.")
    credentials = service_account.Credentials.from_service_account_info(
        json.loads(service_json), scopes=DRIVE_SCOPES
    )
    service = build("drive", "v3", credentials=credentials, cache_discovery=False)
    return service, file_id


def download_drive_file(file_id):
    service_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not service_json:
        raise RuntimeError("Falta GOOGLE_SERVICE_ACCOUNT_JSON en Render.")
    credentials = service_account.Credentials.from_service_account_info(
        json.loads(service_json), scopes=DRIVE_SCOPES
    )
    service = build("drive", "v3", credentials=credentials, cache_discovery=False)
    metadata = service.files().get(fileId=file_id, fields="id,name,mimeType,modifiedTime").execute()
    if metadata.get("mimeType") == "application/vnd.google-apps.spreadsheet":
        request_download = service.files().export_media(
            fileId=file_id,
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        request_download = service.files().get_media(fileId=file_id)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request_download)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buffer.seek(0)
    return buffer, metadata


def colombia_today():
    return datetime.now(ZoneInfo("America/Bogota")).date()


def get_sheets_service():
    file_id = os.environ.get("PLANNING_GOOGLE_DRIVE_FILE_ID", DEFAULT_PLANNING_DRIVE_FILE_ID)
    service_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not service_json:
        raise RuntimeError("Falta GOOGLE_SERVICE_ACCOUNT_JSON en Render.")
    credentials = service_account.Credentials.from_service_account_info(json.loads(service_json), scopes=DRIVE_SCOPES)
    return build("sheets", "v4", credentials=credentials, cache_discovery=False), file_id


def get_planning_drive_service():
    """Devuelve el cliente de Drive y el ID del XLSX de planeación."""
    file_id = os.environ.get("PLANNING_GOOGLE_DRIVE_FILE_ID", DEFAULT_PLANNING_DRIVE_FILE_ID)
    service_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not service_json:
        raise RuntimeError("Falta GOOGLE_SERVICE_ACCOUNT_JSON en Render.")
    credentials = service_account.Credentials.from_service_account_info(
        json.loads(service_json), scopes=["https://www.googleapis.com/auth/drive"]
    )
    return build("drive", "v3", credentials=credentials, cache_discovery=False), file_id


def planning_update_options():
    return [(subject, clei) for subject in ("Matemáticas", "Biología") for clei in ("CLEI 1", "CLEI 2", "CLEI 3", "CLEI 4", "CLEI 5-6")]


def next_week_number(rows):
    numbers = [r.get("week_number", 0) for r in rows if r.get("week_number")]
    return (max(numbers) + 1) if numbers else 1


def build_weekly_proposal(rows, selected):
    current = next_week_number(rows)
    previous = {}
    for row in rows:
        key=(row["subject"], row["group"])
        if row.get("week_number", 0) == current - 1:
            previous[key]=row
    proposal=[]
    for subject, clei in selected:
        prior=previous.get((subject, clei))
        if not prior:
            continue
        proposal.append({
            "subject": subject, "group": clei, "week": f"Semana {current}",
            "date_label": "Por programar", "theme": prior["theme"],
            "objective": prior.get("objective", "No registrado"), "activity": prior.get("activity", "No registrada"),
            "status": "PENDIENTE DE VALIDACIÓN", "source": subject,
        })
    return current, proposal


PLANNING_DECISIONS = {
    "avanzar": "AVANZAR",
    "continuar": "CONTINUAR",
    "reprogramar": "REPROGRAMAR",
}


def curriculum_group_key(value):
    """Convierte CLEI romanos y arábigos a una clave común para filtrar la malla."""
    text = normalize_header(value).replace("-", " ")
    match = re.search(r"(?:clei|nivel|grado|grupo)\s*(vi|iv|iii|ii|i|[1-6])(?:\s|$)", text)
    if not match:
        return text
    raw = match.group(1).upper()
    roman_values = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6}
    number = roman_values.get(raw)
    if number is None and raw.isdigit():
        number = int(raw)
    return f"clei {number}" if number else text


def curriculum_topics(buffer, subject, group, workbook=None):
    """Lee únicamente la malla curricular correspondiente a la materia y al CLEI."""
    if workbook is None:
        from openpyxl import load_workbook
        buffer.seek(0)
        workbook = load_workbook(buffer, data_only=True, read_only=True)

    subject_key = normalize_header(subject)
    subject_aliases = {
        "matematicas": {"mallacurrimat", "mallacurricularmat", "mallamat"},
        "biologia": {"mallacurribio", "mallacurricularbio", "mallabio"},
    }
    sheet_aliases = subject_aliases.get(subject_key, set())
    requested_group = curriculum_group_key(group)
    # El menú agrupa CLEI 5 y CLEI 6, pero la malla debe salir únicamente de
    # la columna CLEI 5.
    if requested_group == "clei 5-6" or normalize_header(group).replace(" ", "") in {"clei5-6", "cleiv-vi"}:
        requested_group = "clei 5"
    # CLEI 1 no corresponde a estas dos materias en la planeación docente.
    # Se muestra explícitamente para que el usuario pueda registrar N/A y no
    # se herede por error el primer tema de la malla.
    if requested_group == "clei 1":
        return ["N/A"]
    topics = []
    seen_topics = set()
    topic_headers = {"tema", "temas", "temacurricular", "temascurriculares", "ej tematico", "ejetematico", "eje tematico", "contenido", "contenidos", "saber", "saberes"}
    group_headers = {"clei", "nivel", "grado", "grupo"}

    def is_topic_header(value):
        normalized = normalize_header(value).strip()
        if normalized in topic_headers:
            return True
        return bool(re.match(r"^(tema|temas|contenido|contenidos|eje tematic|saber|saberes)(?:\s|/|-|:|\(|$)", normalized))

    def is_numeric(value):
        text = normalize_header(value)
        return bool(re.fullmatch(r"(?:p(?:ag|agina)?s?\.?\s*)?\d+(?:[.,]\d+)?(?:\s*(?:[-–/]|a|y)\s*\d+(?:[.,]\d+)?)?", text))

    def header_group_key(value):
        """Acepta encabezados como CLEI I, CLEI 1, I o 1."""
        text = normalize_header(value).strip()
        if re.fullmatch(r"(?:vi|iv|iii|ii|i|[1-6])", text):
            raw = text.upper()
            roman_values = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6}
            number = roman_values.get(raw)
            if number is None and raw.isdigit():
                number = int(raw)
            return f"clei {number}" if number else ""
        return curriculum_group_key(value) if re.search(r"(?:clei|nivel|grado|grupo)", text) else ""

    for worksheet in workbook.worksheets:
        sheet_key = re.sub(r"[^a-z0-9]", "", normalize_header(worksheet.title))
        # No usar coincidencias parciales (por ejemplo, una hoja de clases o de
        # planeación que también contenga la palabra "biología").
        if sheet_key not in sheet_aliases:
            continue

        rows = [[clean_text(value) for value in values] for values in worksheet.iter_rows(values_only=True)]
        rows = [cells for cells in rows if any(cells)]
        if not rows:
            continue

        # Algunas mallas tienen una fila de encabezados (por ejemplo, la fila
        # 14) con un CLEI por columna y los temas debajo de cada columna:
        #     CLEI I | CLEI II | CLEI III ...
        #     Tema 1  | Tema 1   | Tema 1    ...
        # En ese formato no hay una columna "CLEI" para filtrar por fila.
        horizontal_candidates = []
        for row_index, cells in enumerate(rows[:50]):
            columns_by_group = {}
            for column, cell in enumerate(cells):
                group_key = header_group_key(cell)
                if group_key:
                    columns_by_group[column] = group_key
            if len(set(columns_by_group.values())) >= 2:
                has_clei_label = any("clei" in normalize_header(cell) for cell in cells)
                horizontal_candidates.append((has_clei_label, len(columns_by_group), row_index, columns_by_group))

        horizontal_header = None
        if horizontal_candidates:
            # Se prioriza una fila que diga explícitamente CLEI; si la malla
            # solo usa 1, 2, 3..., se toma la fila con más columnas numeradas.
            _, _, header_row_index, columns_by_group = max(horizontal_candidates, key=lambda item: (item[0], item[1], -item[2]))
            horizontal_header = (header_row_index, columns_by_group)

        if horizontal_header:
            header_row_index, columns_by_group = horizontal_header
            for column, column_group in columns_by_group.items():
                if column_group != requested_group:
                    continue
                for cells in rows[header_row_index + 1:]:
                    if column >= len(cells):
                        continue
                    topic = cells[column]
                    if not topic or is_numeric(topic) or is_topic_header(topic):
                        continue
                    topic_key = normalize_header(topic)
                    if topic_key not in seen_topics:
                        seen_topics.add(topic_key)
                        topics.append(topic)
            continue

        # Primero se localiza la cabecera. Si no existe una cabecera estándar,
        # se elige la columna con más texto y se descartan columnas numéricas
        # (consecutivos, códigos o indicadores de la malla).
        header_topic_index = None
        header_group_index = None
        header_row_index = None
        for row_index, cells in enumerate(rows[:50]):
            normalized = [normalize_header(value) for value in cells]
            topic_index = next((i for i, value in enumerate(normalized) if is_topic_header(value)), None)
            if topic_index is not None:
                header_topic_index = topic_index
                header_group_index = next((i for i, value in enumerate(normalized) if value in group_headers), None)
                header_row_index = row_index
                break

        if header_topic_index is None:
            max_columns = max(len(row) for row in rows)
            scores = []
            for column in range(max_columns):
                values = [row[column] for row in rows if column < len(row) and row[column]]
                text_values = [value for value in values if not is_numeric(value)]
                scores.append((len(text_values), sum(len(value) for value in text_values), column))
            header_topic_index = max(scores)[2] if scores else 0
            header_row_index = -1

        for cells in rows[header_row_index + 1:]:
            topic = cells[header_topic_index] if header_topic_index < len(cells) else ""
            if not topic or is_numeric(topic) or is_topic_header(topic):
                continue

            if header_group_index is not None and header_group_index < len(cells):
                row_group = curriculum_group_key(cells[header_group_index])
                if row_group and row_group != requested_group:
                    continue

            topic_key = normalize_header(topic)
            if topic_key not in seen_topics:
                seen_topics.add(topic_key)
                topics.append(topic)
    return topics


def curriculum_topic_pages(buffer, subject, group, topics=None, workbook=None):
    """Obtiene la página que está inmediatamente a la derecha de cada tema.

    La malla usa una distribución horizontal por CLEI. Esta lectura conserva la
    función de temas existente y busca la página en la celda contigua, aceptando
    valores como ``12``, ``12-13`` o ``Pág. 12``.
    """
    from openpyxl import load_workbook

    if workbook is None:
        buffer.seek(0)
        workbook = load_workbook(buffer, data_only=True, read_only=True)
    topics = topics if topics is not None else curriculum_topics(buffer, subject, group, workbook=workbook)
    topic_keys = {normalize_header(topic): topic for topic in topics if clean_text(topic) and clean_text(topic).upper() != "N/A"}
    pages = {topic: "" for topic in topics}
    if not topic_keys:
        return pages

    subject_key = normalize_header(subject)
    subject_aliases = {
        "matematicas": {"mallacurrimat", "mallacurricularmat", "mallamat"},
        "biologia": {"mallacurribio", "mallacurricularbio", "mallabio"},
    }
    sheet_aliases = subject_aliases.get(subject_key, set())

    def is_page_value(value):
        text = clean_text(value)
        normalized = normalize_header(text)
        return bool(re.fullmatch(r"(?:p(?:ag|agina)?s?\.?\s*)?\d+(?:\.0)?(?:\s*(?:[-–/]|a|y)\s*\d+(?:\.0)?)?", normalized))

    def normalize_page(value):
        text = clean_text(value)
        # Excel puede entregar un número entero como 77.0 cuando la celda
        # tiene formato numérico. En la planeación debe verse simplemente 77.
        text = re.sub(r"(?<=\d)\.0(?=\s|$)", "", text)
        return text

    for worksheet in workbook.worksheets:
        sheet_key = re.sub(r"[^a-z0-9]", "", normalize_header(worksheet.title))
        if sheet_key not in sheet_aliases:
            continue
        rows = [[clean_text(value) for value in values] for values in worksheet.iter_rows(values_only=True)]
        for row_index, cells in enumerate(rows):
            for column, cell in enumerate(cells):
                topic = topic_keys.get(normalize_header(cell))
                if not topic:
                    continue

                # Formato habitual: el número de página está inmediatamente
                # a la derecha del tema. También se aceptan hasta tres celdas
                # siguientes porque algunas mallas dejan columnas vacías.
                for candidate in cells[column + 1:column + 4]:
                    if is_page_value(candidate):
                        pages[topic] = normalize_page(candidate)
                        break
                if pages[topic]:
                    continue

                # Algunos archivos colocan la página debajo del tema o usan
                # una celda combinada. Revisamos la fila siguiente y la celda
                # contigua para no perder el vínculo tema-página.
                for next_row in rows[row_index + 1:row_index + 3]:
                    for candidate_column in (column, column + 1, column + 2):
                        if candidate_column < len(next_row) and is_page_value(next_row[candidate_column]):
                            pages[topic] = normalize_page(next_row[candidate_column])
                            break
                    if pages[topic]:
                        break
        break
    return pages


def next_curriculum_topic(topics, current_theme):
    if not topics:
        return ""
    current_key = normalize_header(current_theme)
    for index, topic in enumerate(topics):
        topic_key = normalize_header(topic)
        if topic_key == current_key or current_key in topic_key or topic_key in current_key:
            return topics[index + 1] if index + 1 < len(topics) else ""
    return topics[0]


def append_planning_rows(rows):
    service, file_id = get_sheets_service()
    by_subject = {"Matemáticas": "Matemáticas", "Biología": "Biología"}
    written=[]
    for row in rows:
        sheet = by_subject[row["subject"]]
        values=[["", row["week"], row["date_label"], row["group"], row["theme"], "", row["objective"], row["activity"], row["status"]]]
        service.spreadsheets().values().append(
            spreadsheetId=file_id, range=f"'{sheet}'!A:I", valueInputOption="USER_ENTERED", insertDataOption="INSERT_ROWS", body={"values": values}
        ).execute()
        written.append(f"{sheet} / {row['group']}")
    return written


def download_excel_from_drive():
    """Fuente base: estudiantes, asistencia y grupos. No se reemplaza."""
    _, file_id = get_drive_service()
    return download_drive_file(file_id)


def download_planning_from_drive():
    """Fuente adicional: planeación, clases, materiales y análisis de planeación."""
    file_id = os.environ.get("PLANNING_GOOGLE_DRIVE_FILE_ID", DEFAULT_PLANNING_DRIVE_FILE_ID)
    return download_drive_file(file_id)


def download_external_subjects():
    """Consulta solo valores de las dos planillas externas, sin exportar XLSX."""
    if (
        EXTERNAL_SUBJECTS_CACHE["records"] is not None
        and time.monotonic() - EXTERNAL_SUBJECTS_CACHE["loaded_at"] < EXTERNAL_SUBJECTS_CACHE_SECONDS
    ):
        return EXTERNAL_SUBJECTS_CACHE["records"], EXTERNAL_SUBJECTS_CACHE["metadata"]

    sources = [
        ("Alta y CLEI normal", EXTERNAL_HIGH_CLEI_FILE_ID),
        ("Multigrado / Mediana", EXTERNAL_MULTIGRADE_FILE_ID),
    ]
    records = []
    metadata = []
    service_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not service_json:
        raise RuntimeError("Falta GOOGLE_SERVICE_ACCOUNT_JSON en Render.")
    credentials = service_account.Credentials.from_service_account_info(
        json.loads(service_json), scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"]
    )
    service = build("sheets", "v4", credentials=credentials, cache_discovery=False)
    for source, file_id in sources:
        spreadsheet = service.spreadsheets().get(
            spreadsheetId=file_id,
            fields="spreadsheetId,properties(title),sheets(properties(title,gridProperties(rowCount,columnCount)))",
        ).execute()
        ranges = []
        sheet_ranges = []
        for sheet in spreadsheet.get("sheets", []):
            properties = sheet.get("properties", {})
            title = properties.get("title", "")
            grid = properties.get("gridProperties", {})
            rows = min(int(grid.get("rowCount", 100)), 250)
            columns = min(int(grid.get("columnCount", 26)), 40)
            end_column = chr(64 + columns) if columns <= 26 else "AN"
            ranges.append(f"'{title}'!A1:{end_column}{rows}")
            sheet_ranges.append((title, rows, end_column))
        batch_response = service.spreadsheets().values().batchGet(
            spreadsheetId=file_id,
            ranges=ranges,
            majorDimension="ROWS",
        ).execute()
        for (title, _, _), value_range in zip(
            sheet_ranges,
            batch_response.get("valueRanges", []),
        ):
            records.extend(parse_values(
                value_range.get("values", []),
                title,
                source,
                default_year=colombia_today().year,
            ))
        metadata.append({"source": source, "name": spreadsheet.get("properties", {}).get("title", "")})
    records = consolidate_records(records)
    # Las planillas conservan encabezados históricos y, ocasionalmente,
    # fechas futuras de programación. Otras materias debe mostrar únicamente
    # clases del ciclo académico actual que ya pudieron dictarse.
    today = colombia_today()
    records = [
        item for item in records
        if item.get("class_date")
        and CYCLE_START <= item["class_date"] <= today
    ]
    EXTERNAL_SUBJECTS_CACHE["records"] = records
    EXTERNAL_SUBJECTS_CACHE["metadata"] = metadata
    EXTERNAL_SUBJECTS_CACHE["loaded_at"] = time.monotonic()
    return records, metadata


def upload_planning_to_drive(buffer, metadata=None, created=None):
    """Actualiza el archivo de planeación, Excel o Google Sheets."""
    if metadata and metadata.get("mimeType") == "application/vnd.google-apps.spreadsheet":
        return update_native_google_sheet(created or [])

    service, file_id = get_planning_drive_service()
    buffer.seek(0)
    media = MediaIoBaseUpload(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        resumable=False,
    )
    return service.files().update(
        fileId=file_id,
        media_body=media,
        fields="id,name,modifiedTime",
    ).execute()


def update_native_google_sheet(created):
    """Crea los bloques nuevos directamente en un Google Sheet nativo."""
    if not created:
        return {"id": os.environ.get("PLANNING_GOOGLE_DRIVE_FILE_ID", DEFAULT_PLANNING_DRIVE_FILE_ID)}

    service, file_id = get_sheets_service()
    spreadsheet = service.spreadsheets().get(
        spreadsheetId=file_id,
        fields="sheets(properties(sheetId,title))",
    ).execute()
    sheet_ids = {
        item["properties"]["title"]: item["properties"]["sheetId"]
        for item in spreadsheet.get("sheets", [])
    }
    requests = []
    cleis = ["CLEI I", "CLEI II", "CLEI III", "CLEI IV", "CLEI V", "CLEI VI"]

    for item in created:
        subject = item["subject"]
        sheet_id = sheet_ids.get(subject)
        if sheet_id is None:
            raise RuntimeError(f"No existe la pestaña '{subject}' en el Google Sheet.")
        target_start, target_end = [int(value) for value in item["rows"].split("-")]
        target_start_index = target_start - 1
        target_end_index = target_end
        source_start_index = target_start_index - len(cleis)

        # Copia estilos, bordes, formatos y dimensiones visuales del bloque anterior.
        requests.append({
            "copyPaste": {
                "source": {
                    "sheetId": sheet_id,
                    "startRowIndex": source_start_index,
                    "endRowIndex": target_start_index,
                    "startColumnIndex": 1,
                    "endColumnIndex": 9,
                },
                "destination": {
                    "sheetId": sheet_id,
                    "startRowIndex": target_start_index,
                    "endRowIndex": target_end_index,
                    "startColumnIndex": 1,
                    "endColumnIndex": 9,
                },
                "pasteType": "PASTE_ALL",
            }
        })
        # Limpia todos los campos editables, incluido Estado (columna I).
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": target_start_index,
                    "endRowIndex": target_end_index,
                    "startColumnIndex": 4,
                    "endColumnIndex": 9,
                },
                "cell": {},
                "fields": "userEnteredValue",
            }
        })
        # Define semana, rango de fechas y CLEI, y conserva vacíos los demás cuadros.
        requests.append({
            "updateCells": {
                "start": {"sheetId": sheet_id, "rowIndex": target_start_index, "columnIndex": 1},
                "rows": [{"values": [{"userEnteredValue": {"stringValue": f"Semana {item['week']}"}}, {"userEnteredValue": {"stringValue": item["date"]}}]}],
                "fields": "userEnteredValue",
            }
        })
        requests.append({
            "updateCells": {
                "start": {"sheetId": sheet_id, "rowIndex": target_start_index, "columnIndex": 3},
                "rows": [{"values": [{"userEnteredValue": {"stringValue": clei}}]} for clei in cleis],
                "fields": "userEnteredValue",
            }
        })
        requests.append({
            "updateCells": {
                "start": {"sheetId": sheet_id, "rowIndex": target_start_index, "columnIndex": 4},
                "rows": [
                    {"values": [
                        {"userEnteredValue": {"stringValue": cell.get("theme", "")}},
                        {"userEnteredValue": {"stringValue": cell.get("page", "")}},
                        {"userEnteredValue": {"stringValue": cell.get("objective", "")}},
                        {"userEnteredValue": {"stringValue": cell.get("activity", "")}},
                        {"userEnteredValue": {"stringValue": cell.get("status", "")}},
                    ]}
                    for cell in item.get("cells", [])
                ],
                "fields": "userEnteredValue",
            }
        })
        # Replica las celdas fusionadas de semana y fecha en el nuevo bloque.
        requests.extend([
            {"mergeCells": {"range": {"sheetId": sheet_id, "startRowIndex": target_start_index, "endRowIndex": target_end_index, "startColumnIndex": 1, "endColumnIndex": 2}, "mergeType": "MERGE_ALL"}},
            {"mergeCells": {"range": {"sheetId": sheet_id, "startRowIndex": target_start_index, "endRowIndex": target_end_index, "startColumnIndex": 2, "endColumnIndex": 3}, "mergeType": "MERGE_ALL"}},
        ])

    return service.spreadsheets().batchUpdate(
        spreadsheetId=file_id,
        body={"requests": requests},
    ).execute()


def planning_sheet_names(selected_subjects):
    allowed = {"Matemáticas", "Biología"}
    return [subject for subject in selected_subjects if subject in allowed]


def last_planning_block(worksheet):
    """Encuentra el último bloque de seis filas identificado por 'Semana N'."""
    matches = []
    for row_number in range(1, worksheet.max_row + 1):
        value = clean_text(worksheet.cell(row_number, 2).value)
        match = re.search(r"semana\s+(\d+)", value, re.IGNORECASE)
        if match:
            matches.append((int(match.group(1)), row_number))
    if not matches:
        raise RuntimeError(f"No se encontró una fila de semana en la hoja {worksheet.title}.")
    return max(matches, key=lambda item: (item[0], item[1]))


def copy_cell_style(source, target):
    from copy import copy
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


def next_planning_week_range(value):
    """Devuelve el rango lunes-viernes de la semana siguiente."""
    from datetime import timedelta

    months = {
        "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
        "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
        "septiembre": 9, "setiembre": 9, "octubre": 10,
        "noviembre": 11, "diciembre": 12,
    }
    if isinstance(value, datetime):
        source_start = value.date()
    elif isinstance(value, date):
        source_start = value
    else:
        text = clean_text(value).lower()
        del_range_match = re.search(
            r"del\s+(\d{1,2})\s+de\s+([a-záéíóú]+)\s+al\s+"
            r"\d{1,2}\s+de\s+([a-záéíóú]+)(?:\s+de\s+(\d{4}))?",
            text,
        )
        range_match = re.search(
            r"(\d{1,2})\s*(?:al|a|[-–])\s*(\d{1,2})\s+de\s+"
            r"([a-záéíóú]+)\s+de\s+(\d{4})",
            text,
        )
        single_match = re.search(
            r"(\d{1,2})\s+de\s+([a-záéíóú]+)\s+de\s+(\d{4})",
            text,
        )
        if del_range_match:
            day, month_name = int(del_range_match.group(1)), del_range_match.group(2)
            year = int(del_range_match.group(4) or colombia_today().year)
        elif range_match:
            day, month_name, year = int(range_match.group(1)), range_match.group(3), int(range_match.group(4))
        elif single_match:
            day, month_name, year = int(single_match.group(1)), single_match.group(2), int(single_match.group(3))
        else:
            raise ValueError(f"No se pudo interpretar la fecha de planeación: {value}")
        month = months.get(month_name)
        if not month:
            raise ValueError(f"Mes no válido en la fecha de planeación: {value}")
        source_start = date(year, month, day)

    next_monday = source_start - timedelta(days=source_start.weekday()) + timedelta(days=7)
    next_friday = next_monday + timedelta(days=4)
    month_names = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    ]
    if next_monday.month == next_friday.month:
        return f"{next_monday.day} al {next_friday.day} de {month_names[next_monday.month - 1]} de {next_friday.year}"
    return (
        f"{next_monday.day} de {month_names[next_monday.month - 1]} "
        f"al {next_friday.day} de {month_names[next_friday.month - 1]} de {next_friday.year}"
    )


def create_empty_planning_blocks(buffer, selected_subjects, decisions=None, curriculum_buffer=None, selected_cleis=None, selected_themes=None, selected_objectives=None, selected_articulations=None, selected_pages=None):
    """Crea la semana: CLEI seleccionados reciben tema de la malla; los demás heredan la semana anterior."""
    from copy import copy
    from openpyxl import load_workbook

    selected_subjects = planning_sheet_names(selected_subjects)
    decisions = decisions or {}
    selected_cleis = selected_cleis or {}
    selected_themes = selected_themes or {}
    selected_objectives = selected_objectives or {}
    selected_articulations = selected_articulations or {}
    selected_pages = selected_pages or {}
    if not selected_subjects:
        raise ValueError("Selecciona al menos una materia válida.")

    buffer.seek(0)
    workbook = load_workbook(buffer)
    created = []
    already_exists = []
    # El archivo conserva seis filas; CLEI V y CLEI VI reciben el mismo tema
    # cuando el menú agrupa ambos como CLEI 5-6.
    cleis = ["CLEI I", "CLEI II", "CLEI III", "CLEI IV", "CLEI V", "CLEI VI"]

    for subject in selected_subjects:
        if subject not in workbook.sheetnames:
            raise RuntimeError(f"No existe la hoja '{subject}' en el archivo de planeación.")
        worksheet = workbook[subject]
        last_week, source_start = last_planning_block(worksheet)
        source_end = source_start + len(cleis) - 1
        target_start = source_end + 1
        target_end = target_start + len(cleis) - 1
        next_week = last_week + 1

        # Si el bloque siguiente ya contiene la semana, no lo duplica.
        existing_next = clean_text(worksheet.cell(target_start, 2).value)
        if re.fullmatch(rf"Semana\s+{next_week}", existing_next, re.IGNORECASE):
            already_exists.append(f"{subject} / Semana {next_week}")
            continue

        new_cells = []
        # Replica estilos, bordes, alineación, alturas y formatos de B:I.
        # La nueva semana recibe la decisión elegida para cada CLEI.
        for offset, clei in enumerate(cleis):
            source_row = source_start + offset
            target_row = target_start + offset
            prior = {
                "theme": worksheet.cell(source_row, 5).value,
                "observations": worksheet.cell(source_row, 6).value,
                "objective": worksheet.cell(source_row, 7).value,
                "activity": worksheet.cell(source_row, 8).value,
                "status": worksheet.cell(source_row, 9).value,
            }
            worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
            worksheet.row_dimensions[target_row].hidden = worksheet.row_dimensions[source_row].hidden
            for column in range(2, 10):
                copy_cell_style(worksheet.cell(source_row, column), worksheet.cell(target_row, column))
                worksheet.cell(target_row, column).value = None
            worksheet.cell(target_row, 4).value = clei
            grouped_5_6 = clei in {"CLEI V", "CLEI VI"} and "CLEI 5-6" in selected_cleis.get(subject, [])
            if clei in selected_cleis.get(subject, []) or grouped_5_6:
                theme = selected_themes.get((subject, "CLEI 5-6"), "") if grouped_5_6 else selected_themes.get((subject, clei), "")
                objective_key = (subject, "CLEI 5-6") if grouped_5_6 else (subject, clei)
                objective = selected_objectives.get(objective_key, "")
                activity = selected_articulations.get(objective_key, "")
                if not theme:
                    raise ValueError(f"Selecciona un tema de la malla para {subject} / {clei}.")
                page = selected_pages.get(objective_key, "")
                # La columna F del archivo de planeación es una sola casilla
                # destinada a la página; se escribe únicamente su valor.
                observations = page
                status = "TEMA SELECCIONADO DE LA MALLA"
            else:
                theme = prior["theme"]
                observations = prior["observations"]
                objective = prior["objective"]
                activity = prior["activity"]
                status = prior["status"]
            worksheet.cell(target_row, 5).value = theme
            worksheet.cell(target_row, 6).value = observations
            worksheet.cell(target_row, 7).value = objective
            worksheet.cell(target_row, 8).value = activity
            worksheet.cell(target_row, 9).value = status
            new_cells.append({"group": clei, "theme": clean_text(theme), "page": clean_text(observations), "objective": clean_text(objective), "activity": clean_text(activity), "status": clean_text(status), "mode": "malla" if clei in selected_cleis.get(subject, []) or grouped_5_6 else "heredado"})

        # La semana y la fecha ocupan verticalmente todo el bloque, como en el archivo original.
        for merged_range in (f"B{target_start}:B{target_end}", f"C{target_start}:C{target_end}"):
            worksheet.merge_cells(merged_range)
        worksheet.cell(target_start, 2).value = f"Semana {next_week}"
        next_date_range = next_planning_week_range(worksheet.cell(source_start, 3).value)
        worksheet.cell(target_start, 3).value = next_date_range

        created.append({
            "subject": subject,
            "week": next_week,
            "rows": f"{target_start}-{target_end}",
            "date": next_date_range,
            "cells": new_cells,
        })

    buffer_out = io.BytesIO()
    workbook.save(buffer_out)
    buffer_out.seek(0)
    return buffer_out, created, already_exists


def read_planning_rows(buffer):
    from openpyxl import load_workbook

    buffer.seek(0)
    workbook = load_workbook(buffer, data_only=True, read_only=True)
    classes = []
    for worksheet in workbook.worksheets:
        sheet_name = clean_text(worksheet.title)
        normalized_sheet = normalize_header(sheet_name)
        if normalized_sheet not in PLANNING_SHEETS:
            continue
        rows = list(worksheet.iter_rows(values_only=True))
        for row in rows[1:]:
            values = list(row)
            if not any(clean_text(value) for value in values):
                continue
            if normalized_sheet in {"tecnico laboral", "comunidad terapeutica"}:
                group = sheet_name
                subject = clean_text(values[0]) if len(values) > 0 else ""
                class_date = values[1] if len(values) > 1 else None
                theme = clean_text(values[2]) if len(values) > 2 else ""
                observations = clean_text(values[3]) if len(values) > 3 else ""
                week = clean_text(values[4]) if len(values) > 4 else ""
                drive_link = clean_text(values[6]) if len(values) > 6 else ""
            else:
                group = clean_text(values[0]) if len(values) > 0 else ""
                subject = clean_text(values[1]) if len(values) > 1 else ""
                class_date = values[2] if len(values) > 2 else None
                theme = clean_text(values[3]) if len(values) > 3 else ""
                observations = clean_text(values[4]) if len(values) > 4 else ""
                week = clean_text(values[5]) if len(values) > 5 else ""
                drive_link = clean_text(values[7]) if len(values) > 7 else ""
            normalized_date = parse_date(class_date)
            if not normalized_date:
                continue
            context = classify_context(sheet_name, group)
            no_class = is_no_class(theme, observations)
            cycle = cycle_for_date(normalized_date)
            classes.append({
                "date": normalized_date,
                "date_label": format_date(normalized_date),
                "context": context,
                "context_class": {
                    "Multigrado": "multi",
                    "Técnico Laboral": "tecnico",
                    "Comunidad Terapéutica": "terapeutica",
                    "Alta": "alta",
                }[context],
                "group": group or sheet_name,
                "subject": subject or "Sin asignatura",
                "theme": theme or "Sin tema registrado",
                "observations": observations,
                "week": week,
                "cycle": cycle["cycle"] if cycle else None,
                "cycle_week": cycle["week"] if cycle else None,
                "no_class": no_class,
                "novelty": observations if no_class and observations else (theme if no_class else ""),
                "drive_link": drive_link,
                "source": sheet_name,
            })
    classes.sort(key=lambda item: item["date"], reverse=True)
    return classes


def read_additional_planning_rows(buffer):
    """Lee la hoja adicional con pestañas de materia y filas por semana/CLEI."""
    from openpyxl import load_workbook
    buffer.seek(0)
    workbook = load_workbook(buffer, data_only=True, read_only=True)
    planning = []
    months = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
    for worksheet in workbook.worksheets:
        sheet_name = clean_text(worksheet.title)
        subject_key = normalize_header(sheet_name)
        if subject_key not in {"matematicas", "biologia", "ciencias naturales", "cienciasnaturales"}:
            continue
        subject = "Matemáticas" if subject_key == "matematicas" else ("Biología" if subject_key == "biologia" else "Ciencias Naturales")
        current_week = ""
        current_range = ""
        block_rows_seen = 0
        for row in worksheet.iter_rows(values_only=True):
            values = list(row)
            cells = [clean_text(value) for value in values]
            if len(cells) > 1 and re.match(r"^semana\s+\d+", cells[1], re.IGNORECASE):
                current_week = cells[1]
                current_range = cells[2] if len(cells) > 2 else ""
                block_rows_seen = 0
            # Formato real: columna D = CLEI, E = tema, G = objetivo, H = actividad, I = estado.
            group = cells[3] if len(cells) > 3 else ""
            theme = cells[4] if len(cells) > 4 else ""
            objective = cells[6] if len(cells) > 6 else ""
            activity = cells[7] if len(cells) > 7 else ""
            status = cells[8] if len(cells) > 8 else ""
            group = normalize_planning_clei(group)
            if group:
                block_rows_seen += 1
            # Una fila que conserva solo semana/CLEI después de borrar una
            # prueba no es una planeación válida y no debe reaparecer como
            # "Pendiente" en la plataforma.
            if not group or not current_week or block_rows_seen > 6:
                continue
            if not any((theme, objective, activity, status)):
                continue
            if theme.upper() == "N/A":
                theme = "Sin planeación registrada"
            observations = " | ".join(item for item in (objective, activity, status) if item and item.upper() != "N/A")
            planning.append({
                "date": date(2026, 1, 1), "date_label": current_range or "Fecha por definir",
                "context": "Alta", "context_class": "alta", "group": group,
                "subject": subject, "theme": theme or "Sin tema registrado", "observations": observations,
                "objective": objective or "No registrado", "activity": activity or "No registrada", "status": status or "Pendiente por diligenciar",
                "week": current_week, "week_number": int(re.search(r"\d+", current_week).group()) if re.search(r"\d+", current_week) else 999, "cycle": None, "cycle_week": None,
                "no_class": False, "novelty": "", "drive_link": "", "source": sheet_name,
            })
    clei_order = {"CLEI 1": 1, "CLEI 2": 2, "CLEI 3": 3, "CLEI 4": 4, "CLEI 5-6": 5}
    planning.sort(key=lambda item: (item["week_number"], clei_order.get(item["group"], 99), item["subject"]))
    # CLEI V y CLEI VI comparten la misma planeación. Al normalizarlos a CLEI 5-6,
    # consolidamos duplicados para mostrar una sola tarjeta por semana y materia.
    consolidated = []
    seen = {}
    for item in planning:
        if item["group"] != "CLEI 5-6":
            consolidated.append(item)
            continue
        # V y VI representan el mismo CLEI: una sola tarjeta por semana y materia,
        # incluso si una de las dos filas tiene una variación menor en el texto.
        key = (item["week_number"], normalize_header(item["subject"]))
        previous_index = seen.get(key)
        if previous_index is None:
            seen[key] = len(consolidated)
            consolidated.append(item)
        else:
            previous = consolidated[previous_index]
            # Conserva el registro con más información, sin duplicar la tarjeta.
            current_score = len(item.get("objective", "")) + len(item.get("activity", ""))
            previous_score = len(previous.get("objective", "")) + len(previous.get("activity", ""))
            if current_score > previous_score:
                consolidated[previous_index] = item
    return consolidated


def latest_planning_by_subject_group(buffer):
    """Devuelve la última clase dictada por materia y CLEI desde el Excel base."""
    menu_clei = {
        "CLEI 1": "CLEI I",
        "CLEI 2": "CLEI II",
        "CLEI 3": "CLEI III",
        "CLEI 4": "CLEI IV",
        "CLEI 5-6": "CLEI 5-6",
    }
    latest = {}
    for item in read_planning_rows(buffer):
        group_text = " ".join(
            clean_text(item.get(field, ""))
            for field in ("group", "subject", "context", "source")
        )
        normalized_group = normalize_planning_clei(group_text)
        if not normalized_group:
            continue
        display_group = menu_clei.get(normalized_group, item.get("group", ""))
        class_date = item.get("date")
        if not class_date or not display_group:
            continue
        subject_key = normalize_header(item.get("subject", ""))
        if "matematic" in subject_key or subject_key in {"mate", "mates", "fisica"}:
            display_subject = "Matemáticas"
        elif "biolog" in subject_key or "ciencias natural" in subject_key or subject_key in {"ciencias", "ciencia"}:
            display_subject = "Biología"
        else:
            display_subject = item.get("subject", "")
        key = (display_subject, display_group)
        previous = latest.get(key)
        if previous is None or class_date >= previous["date"]:
            combined = f"{item.get('theme', '')} {item.get('observations', '')}"
            normalized = normalize_header(combined)
            if is_no_class(item.get("theme", ""), item.get("observations", "")):
                execution_status = "No se dictó / hubo novedad"
            elif "sin novedad" in normalized or "sin novedades" in normalized:
                execution_status = "Se dictó sin novedad"
            else:
                execution_status = "Registro disponible; revisar novedades"
            latest[key] = {
                "theme": item.get("theme", "") or "Sin tema registrado",
                "date": class_date,
                "date_label": item.get("date_label", "Fecha por definir"),
                "status": item.get("observations", "") or "Sin novedad registrada",
                "execution_status": execution_status,
                "week": item.get("week", "") or "Clase dictada",
            }
    return latest


def planning_examples(buffer, subject, clei, limit=5):
    """Obtiene ejemplos previos de objetivo y actividad para materia/CLEI."""
    from openpyxl import load_workbook

    buffer.seek(0)
    workbook = load_workbook(buffer, data_only=True, read_only=True)
    worksheet = workbook[subject] if subject in workbook.sheetnames else None
    if worksheet is None:
        return []
    examples = []
    requested_group = normalize_planning_clei(clei)
    for values in worksheet.iter_rows(values_only=True):
        cells = [clean_text(value) for value in values]
        if len(cells) < 8:
            continue
        row_group = normalize_planning_clei(cells[3])
        objective = cells[6]
        articulation = cells[7]
        if row_group != requested_group or not objective or not articulation:
            continue
        examples.append({"objetivo": objective, "articulacion_monitor": articulation})
        if len(examples) >= limit:
            break
    return examples


def normalize_planning_clei(value):
    text = clean_text(value).upper().replace("-", " ")
    match = re.search(r"CLEI\s*([IVX]+|[1-6])", text)
    if not match:
        return ""
    raw = match.group(1)
    roman_values = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6}
    number = roman_values[raw] if raw in roman_values else (int(raw) if raw.isdigit() else None)
    if number is None:
        return ""
    if number in {5, 6}:
        return "CLEI 5-6"
    return f"CLEI {number}"


def normalize_header(value):
    text = clean_text(value).lower()
    return "".join(
        character for character in unicodedata.normalize("NFD", text)
        if unicodedata.category(character) != "Mn"
    )


def is_student_sheet(normalized_sheet):
    """Acepta únicamente hojas de notas/asistencia, nunca hojas de control."""
    if "control" in normalized_sheet or "clase" in normalized_sheet:
        return False
    if normalized_sheet.startswith(COMMUNITY_NOTES_SHEET_MARKER):
        return True
    return normalized_sheet in STUDENT_SHEETS


def sheet_context(normalized_sheet, group=""):
    text = f"{normalized_sheet} {normalize_header(group)}"
    if COMMUNITY_NOTES_SHEET_MARKER in text:
        return "Comunidad Terapéutica"
    if normalized_sheet == "mult. asistencia" or "multigrado" in text:
        return "Multigrado"
    return "Alta"


def available_contexts(records=None):
    """Catálogo estable: los filtros deben aparecer aunque una fuente no tenga filas."""
    return CONTEXT_OPTIONS.copy()


def find_column(headers, aliases, fallback=None):
    aliases = [normalize_header(alias) for alias in aliases]
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        if any(alias in normalized for alias in aliases):
            return index
    return fallback


def is_attendance_grades_sheet(headers):
    """Distingue la hoja de notas/asistencia de una hoja de control de clase."""
    normalized = [normalize_header(value) for value in headers if clean_text(value)]
    markers = ("asistencia", "asistio", "nota", "calificacion", "promedio", "observacion")
    return any(any(marker in header for marker in markers) for header in normalized)


def read_student_records(buffer):
    from openpyxl import load_workbook

    buffer.seek(0)
    workbook = load_workbook(buffer, data_only=True, read_only=True)
    records = []
    for worksheet in workbook.worksheets:
        sheet_name = clean_text(worksheet.title)
        normalized_sheet = normalize_header(sheet_name)
        if not is_student_sheet(normalized_sheet):
            continue
        rows = worksheet.iter_rows(values_only=True)
        header = list(next(rows, ()))
        if not is_attendance_grades_sheet(header):
            continue
        # The old fixed positions remain fallbacks, but headers take priority.
        date_col = find_column(header, ["fecha"], 0)
        name_col = find_column(header, ["nombre", "estudiante", "alumno"], 2)
        identification_col = find_column(header, ["identificacion", "documento", "cedula", "doc"], 3)
        clei_col = find_column(header, ["clei", "nivel"], 4)
        group_col = find_column(header, ["grupo"], 5)
        math_attendance_col = find_column(header, ["asistencia matematicas", "asistio matematicas", "matematicas asistencia", "matematicas asistio", "matematicas presente", "mate asistencia"], 8)
        math_grade_col = find_column(header, ["nota matematicas", "calificacion matematicas", "matematicas nota", "matematicas calificacion", "matematicas promedio"], 9)
        science_attendance_col = find_column(header, ["asistencia ciencias", "asistio ciencias", "ciencias asistencia", "ciencias asistio", "ciencias presente", "ciencias naturales asistencia"], 6)
        science_grade_col = find_column(header, ["nota ciencias", "calificacion ciencias", "ciencias nota", "ciencias calificacion", "ciencias naturales nota", "ciencias naturales promedio"], 7)
        observation_col = find_column(header, ["observacion", "observaciones"], 10)
        for row in rows:
            values = list(row)
            if not any(clean_text(value) for value in values):
                continue
            get = lambda index: values[index] if index is not None and index < len(values) else ""
            name = clean_text(get(name_col))
            identification = clean_text(get(identification_col))
            group = clean_text(get(group_col)) or sheet_name
            if not name and not identification:
                continue
            # Algunas pestañas de notas incluyen filas-resumen o filas de
            # asignatura dentro de la tabla; no son estudiantes.
            if normalize_header(name) in {"matematicas", "biologia", "ciencias naturales", "cienciasnaturales", "tema", "total", "promedio"}:
                continue
            clei_label = {
                "clei 2": "2", "clei 3a": "3A", "clei3b": "3B",
                "clei 4": "4", "clei 5-6": "5-6", "mult. asistencia": "Multigrado",
            }.get(normalized_sheet, clean_text(get(clei_col)))
            class_date = get(date_col)
            records.append({
                "sheet": sheet_name, "name": name, "identification": identification,
                "group": group,
                "context": sheet_context(normalized_sheet, group),
                "date": parse_date(class_date),
                "date_label": format_date(class_date) if parse_date(class_date) else "",
                "clei": clei_label,
                "science_attendance": clean_text(get(science_attendance_col)),
                "science_grade": clean_text(get(science_grade_col)),
                "math_attendance": clean_text(get(math_attendance_col)),
                "math_grade": clean_text(get(math_grade_col)),
                "observation": clean_text(get(observation_col)),
            })
    return records


def detect_probable_exits(records, classes):
    """Detecta posibles salidas: dos o más clases posteriores sin registro del estudiante."""
    class_dates = {}
    for item in classes:
        if item.get("no_class") or not item.get("date"):
            continue
        key = (item.get("group", ""), item.get("context", ""))
        class_dates.setdefault(key, set()).add(item["date"])
    students = {}
    for item in records:
        if not item.get("name") or not item.get("date"):
            continue
        key = (item["name"], item.get("identification", ""), item.get("group", ""))
        current = students.setdefault(key, {
            "name": item["name"], "identification": item.get("identification", ""),
            "group": item.get("group", ""), "context": item.get("context", ""),
            "clei": item.get("clei", ""), "last_date": item["date"],
        })
        current["last_date"] = max(current["last_date"], item["date"])
    exits = []
    for student in students.values():
        dates = sorted(date_value for date_value in class_dates.get((student["group"], student["context"]), set()) if date_value > student["last_date"])
        if len(dates) < 2:
            continue
        student["last_date_label"] = format_date(student["last_date"])
        student["latest_class_label"] = format_date(dates[-1])
        student["future_class_count"] = len(dates)
        student["days_since_last"] = (dates[-1] - student["last_date"]).days
        exits.append(student)
    return sorted(exits, key=lambda item: (-item["future_class_count"], item["last_date"], item["name"].lower()))


def build_student_matrix(records, clei_filter="", cycle_filter="", week_filter="", context_filter=""):
    filtered = [
        item for item in records
        if (not clei_filter or item["clei"] == clei_filter)
        and (not context_filter or item["context"] == context_filter)
        and (not cycle_filter or (cycle_for_date(item["date"]) and cycle_for_date(item["date"])["cycle"] == int(cycle_filter)))
        and (not week_filter or (cycle_for_date(item["date"]) and cycle_for_date(item["date"])["week"] == int(week_filter)))
    ]
    date_map = {}
    students = {}
    for item in filtered:
        if not item["date"]:
            continue
        date_key = item["date"].isoformat()
        date_map[date_key] = item["date_label"]
        student_key = (item["name"], item["identification"], item["group"])
        student = students.setdefault(student_key, {
            "name": item["name"],
            "group": item["group"],
            "clei": item["clei"],
            "dates": {},
            "math_grades": [],
            "science_grades": [],
        })
        cell = student["dates"].setdefault(date_key, {
            "science_attendance": "",
            "science_grade": "",
            "math_attendance": "",
            "math_grade": "",
        })
        cell["science_attendance"] = item["science_attendance"] or cell["science_attendance"]
        cell["science_grade"] = item["science_grade"] or cell["science_grade"]
        cell["math_attendance"] = item["math_attendance"] or cell["math_attendance"]
        cell["math_grade"] = item["math_grade"] or cell["math_grade"]
        for field, target in (("science_grade", "science_grades"), ("math_grade", "math_grades")):
            try:
                student[target].append(float(item[field].replace(",", ".")))
            except (AttributeError, TypeError, ValueError):
                pass

    dates = sorted(date_map.items(), reverse=True)
    matrix = []
    for student in sorted(students.values(), key=lambda value: value["name"].lower()):
        student["math_average"] = round(sum(student["math_grades"]) / len(student["math_grades"]), 2) if student["math_grades"] else None
        student["science_average"] = round(sum(student["science_grades"]) / len(student["science_grades"]), 2) if student["science_grades"] else None
        matrix.append(student)
    return matrix, dates


def get_dashboard_data(context_filter=""):
    if not DRIVE_ENABLED:
        return {
            "classes": [],
            "context_filter": context_filter,
            "current_cycle": current_cycle(),
            "total_classes": 0,
            "total_students": "—",
            "total_groups": 0,
            "drive_updated": "",
            "data_error": "La conexión con Google Drive está pausada. La activaremos nuevamente con el archivo nuevo.",
        }
    try:
        buffer, metadata = download_excel_from_drive()
        all_classes = read_planning_rows(buffer)
        student_records = read_student_records(buffer)
        valid_filters = {"Alta", "Multigrado", "Técnico Laboral", "Comunidad Terapéutica"}
        if context_filter not in valid_filters:
            context_filter = ""
        classes = [
            item for item in all_classes
            if not context_filter or item["context"] == context_filter
        ]
        return {
            "classes": classes,
            "context_filter": context_filter,
            "current_cycle": current_cycle(),
            "total_classes": len(all_classes),
            "total_students": len(student_records),
            "total_groups": len({item["group"] for item in all_classes}),
            "drive_updated": metadata.get("modifiedTime", ""),
            "data_error": None,
        }
    except Exception:
        app.logger.exception("No se pudo leer el Excel privado de Google Drive")
        return {
            "classes": [],
            "context_filter": context_filter,
            "current_cycle": current_cycle(),
            "total_classes": 0,
            "total_students": "—",
            "total_groups": 0,
            "drive_updated": "",
            "data_error": "No se pudo actualizar el Excel desde Google Drive. Revisa la configuración de Render.",
        }


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        if username == ADMIN_USER and password == ADMIN_PASSWORD:
            session.clear()
            session["user"] = username
            return redirect(url_for("home"))
        error = "El usuario o la contraseña no son correctos."
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def home():
    context_filter = request.args.get("context", "").strip()
    data = get_dashboard_data(context_filter)
    return render_template("index.html", current_user=session.get("user"), **data)


@app.route("/grupos")
@login_required
def grupos():
    context_filter = request.args.get("contexto", "").strip()
    if not DRIVE_ENABLED:
        return render_template(
            "grupos.html",
            current_user=session.get("user"),
            matrix=[],
            dates=[],
            cleis=["2", "3A", "3B", "4", "5-6", "Multigrado"],
            clei_filter="",
            contexts=CONTEXT_OPTIONS.copy(),
            context_filter=context_filter,
            cycles=[],
            weeks=[1, 2, 3],
            cycle_filter="",
            week_filter="",
            total_records=0,
            data_error="La conexión con Google Drive está pausada. La activaremos nuevamente con el archivo nuevo.",
            drive_updated="",
        )
    try:
        buffer, metadata = download_excel_from_drive()
        records = read_student_records(buffer)
        clei_filter = request.args.get("clei", "").strip()
        cycle_filter = request.args.get("ciclo", "").strip()
        week_filter = request.args.get("semana", "").strip()
        cleis = ["2", "3A", "3B", "4", "5-6", "Multigrado"]
        cycles = sorted({cycle_for_date(item["date"])["cycle"] for item in records if cycle_for_date(item["date"])})
        weeks = [1, 2, 3]
        contexts = available_contexts(records)
        if context_filter not in contexts:
            context_filter = ""
        if clei_filter not in cleis:
            clei_filter = ""
        if cycle_filter not in {str(value) for value in cycles}:
            cycle_filter = ""
        if week_filter not in {str(value) for value in weeks}:
            week_filter = ""
        matrix, dates = build_student_matrix(records, clei_filter, cycle_filter, week_filter, context_filter)
        return render_template(
            "grupos.html",
            current_user=session.get("user"),
            matrix=matrix,
            dates=dates,
            cleis=cleis,
            clei_filter=clei_filter,
            contexts=contexts,
            context_filter=context_filter,
            cycles=cycles,
            weeks=weeks,
            cycle_filter=cycle_filter,
            week_filter=week_filter,
            total_records=len(matrix),
            data_error=None,
            drive_updated=metadata.get("modifiedTime", ""),
        )
    except Exception:
        app.logger.exception("No se pudo leer estudiantes para Mis grupos")
        return render_template(
            "grupos.html",
            current_user=session.get("user"),
            matrix=[],
            dates=[],
            cleis=[],
            clei_filter="",
            contexts=CONTEXT_OPTIONS.copy(),
            context_filter=context_filter,
            cycles=[],
            weeks=[1, 2, 3],
            cycle_filter="",
            week_filter="",
            total_records=0,
            data_error="No se pudo leer el Excel desde Google Drive.",
            drive_updated="",
        )


@app.route("/actualizar-planeacion", methods=["GET", "POST"])
@login_required
def actualizar_planeacion():
    options = [
        {"value": "Matemáticas", "label": "Matemáticas"},
        {"value": "Biología", "label": "Biología"},
    ]
    today=colombia_today()
    friday=today.weekday()==4
    cleis = ["CLEI I", "CLEI II", "CLEI III", "CLEI IV", "CLEI 5-6"]
    page={"options": options, "cleis": cleis, "topics_by_subject": {item["value"]: {clei: [] for clei in cleis} for item in options}, "pages_by_subject": {item["value"]: {clei: {} for clei in cleis} for item in options}, "last_planning": {}, "selected": [], "selected_cleis": {}, "selected_themes": {}, "selected_pages": {}, "selected_objectives": {}, "selected_articulations": {}, "proposals": {}, "created": [], "already_exists": [], "error": None, "message": None, "is_friday": friday, "only_friday": PLANNING_ONLY_FRIDAY}
    try:
        curriculum_buffer = None
        if DRIVE_ENABLED:
            cache_is_fresh = (
                CURRICULUM_CACHE["topics"]
                and time.monotonic() - CURRICULUM_CACHE["loaded_at"] < CURRICULUM_CACHE_SECONDS
            )
            if cache_is_fresh:
                page["topics_by_subject"] = CURRICULUM_CACHE["topics"]
                page["pages_by_subject"] = CURRICULUM_CACHE["pages"] or page["pages_by_subject"]
            else:
                from openpyxl import load_workbook
                curriculum_buffer, curriculum_metadata = download_excel_from_drive()
                modified_time = curriculum_metadata.get("modifiedTime", "")
                curriculum_buffer.seek(0)
                curriculum_workbook = load_workbook(curriculum_buffer, data_only=True, read_only=True)
                for option in options:
                    for clei in cleis:
                        page["topics_by_subject"][option["value"]][clei] = curriculum_topics(curriculum_buffer, option["value"], clei, workbook=curriculum_workbook)
                        page["pages_by_subject"][option["value"]][clei] = curriculum_topic_pages(curriculum_buffer, option["value"], clei, topics=page["topics_by_subject"][option["value"]][clei], workbook=curriculum_workbook)
                CURRICULUM_CACHE["modified_time"] = modified_time
                CURRICULUM_CACHE["topics"] = page["topics_by_subject"]
                CURRICULUM_CACHE["pages"] = page["pages_by_subject"]
                CURRICULUM_CACHE["loaded_at"] = time.monotonic()
            if curriculum_buffer is None:
                curriculum_buffer, _ = download_excel_from_drive()
            page["last_planning"] = latest_planning_by_subject_group(curriculum_buffer)
        if request.method == "POST":
            selected = planning_sheet_names(request.form.getlist("materia"))
            page["selected"] = selected
            def planning_field_key(value):
                return re.sub(r"\s+", "", normalize_header(value))

            selected_cleis = {
                subject: [clei for clei in cleis if clei in request.form.getlist(f"clei__{planning_field_key(subject)}")]
                for subject in selected
            }
            selected_themes = {
                (subject, clei): clean_text(request.form.get(f"tema__{planning_field_key(subject)}__{planning_field_key(clei)}", ""))
                for subject in selected for clei in selected_cleis.get(subject, [])
            }
            selected_pages = {
                (subject, clei): clean_text(
                    page["pages_by_subject"].get(subject, {}).get(clei, {}).get(
                        selected_themes.get((subject, clei), ""), ""
                    )
                )
                for subject in selected for clei in selected_cleis.get(subject, [])
            }
            selected_objectives = {
                (subject, clei): clean_text(request.form.get(f"objetivo__{planning_field_key(subject)}__{planning_field_key(clei)}", ""))
                for subject in selected for clei in selected_cleis.get(subject, [])
            }
            selected_articulations = {
                (subject, clei): clean_text(request.form.get(f"articulacion__{planning_field_key(subject)}__{planning_field_key(clei)}", ""))
                for subject in selected for clei in selected_cleis.get(subject, [])
            }
            page["selected_cleis"] = selected_cleis
            page["selected_themes"] = selected_themes
            page["selected_pages"] = selected_pages
            page["selected_objectives"] = selected_objectives
            page["selected_articulations"] = selected_articulations
            accion = request.form.get("accion", "crear")
            if accion == "generar_ia":
                try:
                    if not selected:
                        raise GeminiPlanningError("Selecciona al menos una materia para generar una propuesta.")
                    planning_buffer, _ = download_planning_from_drive()
                    for subject, subject_cleis in selected_cleis.items():
                        for clei in subject_cleis:
                            theme = selected_themes.get((subject, clei), "")
                            if not theme:
                                raise GeminiPlanningError(f"Selecciona primero un tema para {subject} / {clei}.")
                            proposal = generate_planning_proposal(
                                subject,
                                clei,
                                theme,
                                planning_examples(planning_buffer, subject, clei),
                            )
                            page["selected_objectives"][(subject, clei)] = proposal["objetivo"]
                            page["selected_articulations"][(subject, clei)] = proposal["articulacion_monitor"]
                            page["proposals"][(subject, clei)] = proposal
                    page["message"] = "Gemini generó propuestas. Revísalas y edítalas si lo deseas antes de actualizar Drive."
                except GeminiPlanningError as exc:
                    page["error"] = str(exc)
                return render_template("actualizar_planeacion.html", current_user=session.get("user"), **page)
            if not selected:
                page["error"] = "Selecciona al menos una materia para crear la nueva estructura."
            elif PLANNING_ONLY_FRIDAY and not friday:
                page["error"] = "La actualización solo puede confirmarse los viernes."
            else:
                buffer, metadata = download_planning_from_drive()
                for subject, subject_cleis in selected_cleis.items():
                    for clei in subject_cleis:
                        theme = selected_themes.get((subject, clei), "")
                        if theme not in page["topics_by_subject"].get(subject, {}).get(clei, []):
                            raise ValueError(f"El tema seleccionado para {subject} / {clei} no pertenece a la malla cargada.")
                updated_buffer, created, already_exists = create_empty_planning_blocks(buffer, selected, selected_cleis=selected_cleis, selected_themes=selected_themes, selected_objectives=selected_objectives, selected_articulations=selected_articulations, selected_pages=selected_pages)
                if created:
                    metadata = upload_planning_to_drive(updated_buffer, metadata=metadata, created=created)
                    page["created"] = created
                    page["already_exists"] = already_exists
                    page["message"] = f"Se actualizó la planeación en {len(created)} hoja(s). Los CLEI seleccionados tomaron tema de la malla y los demás heredaron la semana anterior."
                else:
                    page["already_exists"] = already_exists
                    page["message"] = "La siguiente semana ya estaba creada; no se duplicaron filas."
        return render_template("actualizar_planeacion.html", current_user=session.get("user"), **page)
    except Exception as exc:
        app.logger.exception("No se pudo preparar la actualización semanal")
        page["error"]="No se pudo preparar la actualización. Verifica la conexión y los permisos de edición del archivo de planeación."
        return render_template("actualizar_planeacion.html", current_user=session.get("user"), **page)


@app.route("/estudiantes")
@login_required
def estudiantes():
    page_data = {
        "students": [],
        "cleis": ["2", "3A", "3B", "4", "5-6", "Multigrado"],
        "contexts": CONTEXT_OPTIONS.copy(),
        "clei_filter": "",
        "context_filter": "",
        "search": "",
        "total_students": 0,
        "probable_exits": [],
        "data_error": None,
    }
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("estudiantes.html", current_user=session.get("user"), **page_data)
    try:
        buffer, metadata = download_excel_from_drive()
        records = read_student_records(buffer)
        classes = read_planning_rows(buffer)
        search = request.args.get("buscar", "").strip().lower()
        clei_filter = request.args.get("clei", "").strip()
        context_filter = request.args.get("contexto", "").strip()
        cleis = ["2", "3A", "3B", "4", "5-6", "Multigrado"]
        contexts = available_contexts(records)
        if clei_filter not in cleis:
            clei_filter = ""
        if context_filter not in contexts:
            context_filter = ""
        unique = {}
        for item in records:
            if search and search not in f"{item['name']} {item['group']}".lower():
                continue
            if clei_filter and item["clei"] != clei_filter:
                continue
            if context_filter and item["context"] != context_filter:
                continue
            key = (item["name"], item["identification"], item["group"])
            student = unique.setdefault(key, {
                "name": item["name"], "identification": item["identification"],
                "group": item["group"], "clei": item["clei"], "context": item["context"],
                "dates": set(), "math_grades": [], "science_grades": [], "records": [],
            })
            if item["date"]:
                student["dates"].add(item["date"])
            student["records"].append({
                "date": item["date_label"] or "Sin fecha",
                "math_attendance": item["math_attendance"], "math_grade": item["math_grade"],
                "science_attendance": item["science_attendance"], "science_grade": item["science_grade"],
                "observation": item["observation"],
            })
            for field, target in (("math_grade", "math_grades"), ("science_grade", "science_grades")):
                try:
                    student[target].append(float(item[field].replace(",", ".")))
                except (AttributeError, TypeError, ValueError):
                    pass
        students = []
        for student in unique.values():
            student["date_count"] = len(student.pop("dates"))
            student["math_average"] = round(sum(student["math_grades"]) / len(student["math_grades"]), 2) if student["math_grades"] else None
            student["science_average"] = round(sum(student["science_grades"]) / len(student["science_grades"]), 2) if student["science_grades"] else None
            students.append(student)
        probable_exits = [item for item in detect_probable_exits(records, classes)
                          if (not search or search in f"{item['name']} {item['group']}".lower())
                          and (not clei_filter or item["clei"] == clei_filter)
                          and (not context_filter or item["context"] == context_filter)]
        page_data.update({
            "students": sorted(students, key=lambda item: item["name"].lower()),
            "cleis": cleis, "contexts": contexts, "clei_filter": clei_filter,
            "context_filter": context_filter, "search": request.args.get("buscar", "").strip(),
            "total_students": len(students), "probable_exits": probable_exits,
            "drive_updated": metadata.get("modifiedTime", ""),
        })
        return render_template("estudiantes.html", current_user=session.get("user"), **page_data)
    except Exception:
        app.logger.exception("No se pudo leer estudiantes")
        page_data["data_error"] = "No se pudo leer el Excel desde Google Drive."
        return render_template("estudiantes.html", current_user=session.get("user"), **page_data)



def attendance_value(value):
    text = clean_text(value).lower()
    if text in {"si", "sí", "s", "asistio", "asistió", "presente", "p"}:
        return "Asistió"
    if text in {"no", "n", "ausente", "inasistente", "i"}:
        return "No asistió"
    return clean_text(value) or "Sin registro"


@app.route("/asistencia")
@login_required
def asistencia():
    page_data = {
        "rows": [], "cleis": ["2", "3A", "3B", "4", "5-6", "Multigrado"], "contexts": CONTEXT_OPTIONS.copy(), "context_filter": "",
        "cycles": [], "weeks": [1, 2, 3],
        "clei_filter": [], "cycle_filter": [], "week_filter": [],
        "start_date": "", "end_date": "", "total_present": 0, "total_absent": 0, "total_rows": 0,
        "attendance_stats": [], "total_never_attended": 0,
        "data_error": None,
    }
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("asistencia.html", current_user=session.get("user"), **page_data)
    try:
        buffer, metadata = download_excel_from_drive()
        records = read_student_records(buffer)
        clei_filter = request.args.getlist("clei")
        cycle_filter = request.args.getlist("ciclo")
        week_filter = request.args.getlist("semana")
        context_filter = request.args.get("contexto", "").strip()
        start_date_text = request.args.get("desde", "").strip()
        end_date_text = request.args.get("hasta", "").strip()
        try:
            start_date = datetime.strptime(start_date_text, "%Y-%m-%d").date() if start_date_text else None
        except ValueError:
            start_date_text, start_date = "", None
        try:
            end_date = datetime.strptime(end_date_text, "%Y-%m-%d").date() if end_date_text else None
        except ValueError:
            end_date_text, end_date = "", None
        if start_date and end_date and start_date > end_date:
            start_date, end_date = end_date, start_date
            start_date_text, end_date_text = end_date_text, start_date_text
        cleis = ["2", "3A", "3B", "4", "5-6", "Multigrado"]
        cycles = sorted({cycle_for_date(item["date"])["cycle"] for item in records if cycle_for_date(item["date"])})
        weeks = [1, 2, 3]
        contexts = available_contexts(records)
        if context_filter not in contexts:
            context_filter = ""
        clei_filter = [value for value in clei_filter if value in cleis]
        cycle_filter = [value for value in cycle_filter if value in {str(item) for item in cycles}]
        week_filter = [value for value in week_filter if value in {str(item) for item in weeks}]

        rows = []
        attendance_by_student = {}
        for item in records:
            cycle = cycle_for_date(item["date"])
            if clei_filter and item["clei"] not in clei_filter:
                continue
            if context_filter and item["context"] != context_filter:
                continue
            if cycle_filter and (not cycle or str(cycle["cycle"]) not in cycle_filter):
                continue
            if week_filter and (not cycle or str(cycle["week"]) not in week_filter):
                continue
            if start_date and (not item["date"] or item["date"] < start_date):
                continue
            if end_date and (not item["date"] or item["date"] > end_date):
                continue

            math_status = attendance_value(item["math_attendance"]) if (item["math_attendance"] or item["math_grade"]) else "Sin registro"
            science_status = attendance_value(item["science_attendance"]) if (item["science_attendance"] or item["science_grade"]) else "Sin registro"
            if math_status == "Sin registro" and science_status == "Sin registro":
                continue
            rows.append({
                "name": item["name"], "group": item["group"], "clei": item["clei"],
                "context": item["context"], "date_label": item["date_label"], "date": item["date"],
                "math_status": math_status, "science_status": science_status,
                "observation": clean_text(item.get("observation", "")),
                "cycle": cycle["cycle"] if cycle else "—", "week": cycle["week"] if cycle else "—",
            })

            key = (item["name"], item["identification"], item["group"])
            summary = attendance_by_student.setdefault(key, {
                "name": item["name"], "group": item["group"], "clei": item["clei"], "context": item["context"],
                "math_sessions": 0, "math_present": 0, "math_absent": 0,
                "science_sessions": 0, "science_present": 0, "science_absent": 0,
                "observations": [], "absence_observations": [],
            })
            observation = clean_text(item.get("observation", ""))
            if observation and observation not in summary["observations"]:
                summary["observations"].append(observation)
            for prefix, raw_status, grade in (
                ("math", item["math_attendance"], item["math_grade"]),
                ("science", item["science_attendance"], item["science_grade"]),
            ):
                if not (raw_status or grade):
                    continue
                status = attendance_value(raw_status)
                if observation and status == "No asistió" and observation not in summary["absence_observations"]:
                    summary["absence_observations"].append(observation)
                summary[f"{prefix}_sessions"] += 1
                if status == "Asistió":
                    summary[f"{prefix}_present"] += 1
                elif status == "No asistió":
                    summary[f"{prefix}_absent"] += 1

        rows.sort(key=lambda item: (item["date"] or date.min, item["name"].lower()), reverse=True)
        attendance_stats = []
        for summary in attendance_by_student.values():
            summary["total_sessions"] = summary["math_sessions"] + summary["science_sessions"]
            summary["total_present"] = summary["math_present"] + summary["science_present"]
            summary["total_absent"] = summary["math_absent"] + summary["science_absent"]
            summary["math_rate"] = round(summary["math_present"] * 100 / summary["math_sessions"], 1) if summary["math_sessions"] else 0
            summary["science_rate"] = round(summary["science_present"] * 100 / summary["science_sessions"], 1) if summary["science_sessions"] else 0
            summary["total_rate"] = round(summary["total_present"] * 100 / summary["total_sessions"], 1) if summary["total_sessions"] else 0
            summary["never_attended"] = summary["total_present"] == 0
            summary["possible_cause"] = " · ".join(summary["absence_observations"]) if summary["absence_observations"] else "No hay observación registrada para explicar la inasistencia."
            attendance_stats.append(summary)
        attendance_stats.sort(key=lambda item: (not item["never_attended"], -item["total_absent"], item["total_present"], item["name"].lower()))
        page_data.update({
            "rows": rows, "cleis": cleis, "contexts": contexts, "context_filter": context_filter, "cycles": cycles, "weeks": weeks,
            "clei_filter": clei_filter, "cycle_filter": cycle_filter, "week_filter": week_filter,
            "start_date": start_date_text, "end_date": end_date_text,
            "total_present": sum(row["math_status"] == "Asistió" for row in rows) + sum(row["science_status"] == "Asistió" for row in rows),
            "total_absent": sum(row["math_status"] == "No asistió" for row in rows) + sum(row["science_status"] == "No asistió" for row in rows),
            "total_rows": len(rows), "data_error": None,
            "attendance_stats": attendance_stats,
            "total_never_attended": sum(item["never_attended"] for item in attendance_stats),
            "drive_updated": metadata.get("modifiedTime", ""),
        })
        return render_template("asistencia.html", current_user=session.get("user"), **page_data)
    except Exception:
        app.logger.exception("No se pudo leer asistencia")
        page_data["data_error"] = "No se pudo leer el Excel desde Google Drive."
        return render_template("asistencia.html", current_user=session.get("user"), **page_data)


@app.route("/materiales")
@login_required
def materiales():
    page_data = {
        "materials": [],
        "contexts": ["Alta", "Multigrado", "Técnico Laboral", "Comunidad Terapéutica"],
        "context_filter": "",
        "group_filter": "",
        "search": "",
        "groups": [],
        "data_error": None,
    }
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("materiales.html", current_user=session.get("user"), **page_data)
    try:
        buffer, metadata = download_excel_from_drive()
        classes = read_planning_rows(buffer)
        context_filter = request.args.get("contexto", "").strip()
        group_filter = request.args.get("grupo", "").strip()
        search = request.args.get("buscar", "").strip()
        contexts = CONTEXT_OPTIONS.copy()
        if context_filter not in contexts:
            context_filter = ""
        groups = sorted({item["group"] for item in classes if item["group"]})
        if group_filter not in groups:
            group_filter = ""
        search_lower = search.lower()
        materials = [
            item for item in classes
            if item["drive_link"]
            and (not context_filter or item["context"] == context_filter)
            and (not group_filter or item["group"] == group_filter)
            and (not search_lower or search_lower in f"{item['subject']} {item['theme']} {item['group']} {item['observations']}".lower())
        ]
        page_data.update({
            "materials": materials,
            "context_filter": context_filter,
            "group_filter": group_filter,
            "search": search,
            "groups": groups,
            "drive_updated": metadata.get("modifiedTime", ""),
        })
        return render_template("materiales.html", current_user=session.get("user"), **page_data)
    except Exception:
        app.logger.exception("No se pudieron leer los materiales")
        page_data["data_error"] = "No se pudo leer el Excel desde Google Drive."
        return render_template("materiales.html", current_user=session.get("user"), **page_data)


@app.route("/otras-materias")
@login_required
def otras_materias():
    """Consulta detallada de materias de otras docentes, sin escribir en Drive."""
    page_data = {
        "records": [],
        "matrix": [],
        "dates": [],
        "date_options": [],
        "sources": ["Alta y CLEI normal", "Multigrado / Mediana"],
        "subjects": [],
        "cleis": [],
        "teachers": [],
        "contexts": [],
        "source_filter": request.args.get("fuente", "").strip(),
        "subject_filter": request.args.get("materia", "").strip(),
        "teacher_filter": request.args.get("profesora", "").strip(),
        "context_filter": request.args.get("contexto", "").strip(),
        "date_filter": request.args.get("fecha", "").strip(),
        "search": request.args.get("buscar", "").strip(),
        "only_mismatches": request.args.get("recuperaciones", "") == "1",
        "summary": summarize([]),
        "data_error": None,
    }
    if not EXTERNAL_SUBJECTS_ENABLED:
        page_data["data_error"] = "La consulta de otras materias está temporalmente pausada para proteger la estabilidad de Render. Activa EXTERNAL_SUBJECTS_ENABLED=true cuando el servicio esté estable."
        return render_template("otras_materias.html", current_user=session.get("user"), **page_data)
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("otras_materias.html", current_user=session.get("user"), **page_data)
    try:
        all_records, metadata = download_external_subjects()
        mark_week_mismatches(all_records, cycle_for_date)
        page_data["subjects"] = sorted({item["subject"] for item in all_records})
        page_data["cleis"] = sorted({item["clei"] for item in all_records})
        page_data["teachers"] = sorted({item["teacher"] for item in all_records})
        page_data["contexts"] = sorted({item["context"] for item in all_records})
        page_data["date_options"] = sorted({
            (item["class_date"].isoformat(), item["class_date_label"])
            for item in all_records if item.get("class_date")
        })
        needle = normalize_header(page_data["search"])
        records = [
            item for item in all_records
            if (not page_data["source_filter"] or item["source"] == page_data["source_filter"])
            and (not page_data["subject_filter"] or item["subject"] == page_data["subject_filter"])
            and (not page_data["teacher_filter"] or item["teacher"] == page_data["teacher_filter"])
            and (not page_data["context_filter"] or item["context"] == page_data["context_filter"])
            and (not page_data["date_filter"] or (item.get("class_date") and item["class_date"].isoformat() == page_data["date_filter"]))
            and (not page_data["only_mismatches"] or item["week_mismatch"])
            and (not needle or needle in normalize_header(" ".join(str(item.get(key, "")) for key in ("student", "identification", "group", "subject", "teacher", "clei"))))
        ]
        records.sort(key=lambda item: (item.get("class_date") or date.min, item["student"]), reverse=True)
        matrix, dates = build_external_matrix(
            records,
            clei_filter=request.args.get("clei", "").strip(),
        )
        page_data.update({
            "records": records,
            "matrix": matrix,
            "dates": dates,
            "summary": summarize(records),
            "drive_updated": "; ".join(f"{item['source']}: {item.get('modifiedTime', '')}" for item in metadata),
        })
        return render_template("otras_materias.html", current_user=session.get("user"), **page_data)
    except Exception as exc:
        app.logger.exception("No se pudieron leer las planillas externas de otras materias")
        # Mostrar únicamente el tipo y el texto del error de Google; nunca se
        # incluye el JSON de credenciales ni tokens en la respuesta.
        detail = f"{type(exc).__name__}: {str(exc)}".replace("\n", " ")[:500]
        page_data["data_error"] = f"Error real de Google Sheets: {detail}"
        return render_template("otras_materias.html", current_user=session.get("user"), **page_data)


@app.route("/seguimiento")
@login_required
def seguimiento():
    page_data = {
        "stats": {
            "classes": 0, "students": 0, "groups": 0, "attendance_rate": 0,
            "math_rate": 0, "science_rate": 0, "math_average": None, "science_average": None,
            "risk_count": 0, "low_performance_count": 0, "students_with_grades": 0,
            "records": 0, "attendance_coverage": 0, "recent_rate": 0, "trend": 0,
            "chronic_absence_count": 0, "missing_attendance": 0, "total_sessions": 0,
            "total_present": 0, "total_absent": 0, "math_sessions": 0, "science_sessions": 0,
        },
        "chart_data": json.dumps({"dates": [], "math": [], "science": [], "contexts": [], "risks": []}),
        "interpretations": [], "risk_students": [], "probable_exits": [], "data_error": None,
    }
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("seguimiento.html", current_user=session.get("user"), **page_data)
    try:
        buffer, metadata = download_excel_from_drive()
        classes = read_planning_rows(buffer)
        records = read_student_records(buffer)
        valid_records = [item for item in records if item["date"]]
        probable_exits = detect_probable_exits(valid_records, classes)
        unique_students = {(item["name"], item["identification"], item["group"]) for item in valid_records if item["name"]}
        attendance = {"math": {"sessions": 0, "present": 0, "absent": 0}, "science": {"sessions": 0, "present": 0, "absent": 0}}
        grades = {"math": [], "science": []}
        dates = {}
        student_stats = {}
        records_with_attendance = 0
        missing_attendance = 0
        for item in valid_records:
            student_key = (item["name"], item["identification"], item["group"])
            student = student_stats.setdefault(student_key, {"name": item["name"], "group": item["group"], "clei": item["clei"], "present": 0, "absent": 0, "sessions": 0})
            day = item["date"].isoformat()
            date_entry = dates.setdefault(day, {"label": item["date"].strftime("%d/%m/%Y"), "math_present": 0, "math_absent": 0, "science_present": 0, "science_absent": 0})
            record_has_attendance = False
            for prefix, attendance_value_raw, grade_raw in (("math", item["math_attendance"], item["math_grade"]), ("science", item["science_attendance"], item["science_grade"])):
                if not (attendance_value_raw or grade_raw):
                    continue
                status = attendance_value(attendance_value_raw)
                attendance[prefix]["sessions"] += 1
                student["sessions"] += 1
                if status in {"Asistió", "No asistió"}:
                    record_has_attendance = True
                if status == "Asistió":
                    attendance[prefix]["present"] += 1
                    student["present"] += 1
                    date_entry[f"{prefix}_present"] += 1
                elif status == "No asistió":
                    attendance[prefix]["absent"] += 1
                    student["absent"] += 1
                    date_entry[f"{prefix}_absent"] += 1
                try:
                    grades[prefix].append(float(grade_raw.replace(",", ".")))
                except (AttributeError, TypeError, ValueError):
                    pass
            if record_has_attendance:
                records_with_attendance += 1
            else:
                missing_attendance += 1
        def rate(data):
            return round(data["present"] * 100 / data["sessions"], 1) if data["sessions"] else 0
        math_rate, science_rate = rate(attendance["math"]), rate(attendance["science"])
        total_sessions = attendance["math"]["sessions"] + attendance["science"]["sessions"]
        total_present = attendance["math"]["present"] + attendance["science"]["present"]
        overall_rate = round(total_present * 100 / total_sessions, 1) if total_sessions else 0
        risk_students = []
        low_performance_students = 0
        chronic_absence_count = 0
        for item in student_stats.values():
            item["rate"] = round(item["present"] * 100 / item["sessions"], 1) if item["sessions"] else 0
            item["risk_reason"] = ""
            if item["sessions"] and (item["rate"] < 70 or item["absent"] > item["present"]):
                reasons = []
                if item["rate"] < 70:
                    reasons.append("asistencia menor al 70%")
                if item["absent"] > item["present"]:
                    reasons.append("más inasistencias que asistencias")
                item["risk_reason"] = " y ".join(reasons)
                risk_students.append(item)
            if item["absent"] >= 2:
                chronic_absence_count += 1
        student_grade_averages = {}
        for item in valid_records:
            key = (item["name"], item["identification"], item["group"])
            for prefix in ("math", "science"):
                try:
                    student_grade_averages.setdefault(key, {}).setdefault(prefix, []).append(float(item[f"{prefix}_grade"].replace(",", ".")))
                except (AttributeError, TypeError, ValueError):
                    pass
        for values in student_grade_averages.values():
            all_grades = [grade for subject_grades in values.values() for grade in subject_grades]
            if all_grades and sum(all_grades) / len(all_grades) < 3:
                low_performance_students += 1
        risk_students.sort(key=lambda item: (-item["absent"], item["rate"], item["name"].lower()))
        risk_students = risk_students[:12]
        date_items = sorted(dates.items())
        date_rates = []
        for _, entry in date_items:
            present = sum(entry[f"{prefix}_present"] for prefix in ("math", "science"))
            absent = sum(entry[f"{prefix}_absent"] for prefix in ("math", "science"))
            if present + absent:
                date_rates.append(round(present * 100 / (present + absent), 1))
        recent_rate = date_rates[-4:] and round(sum(date_rates[-4:]) / len(date_rates[-4:]), 1) or 0
        previous_rates = date_rates[-8:-4]
        previous_rate = round(sum(previous_rates) / len(previous_rates), 1) if previous_rates else recent_rate
        trend = round(recent_rate - previous_rate, 1)
        context_counts = {}
        for item in classes:
            context_counts[item["context"]] = context_counts.get(item["context"], 0) + 1
        interpretations = []
        if total_sessions:
            interpretations.append(f"La asistencia global registrada es de {overall_rate}%, calculada sobre {total_sessions} sesiones de Matemáticas y Ciencias Naturales.")
        if math_rate and science_rate:
            better = "Matemáticas" if math_rate >= science_rate else "Ciencias Naturales"
            difference = abs(math_rate - science_rate)
            interpretations.append(f"El mejor comportamiento de asistencia se observa en {better}; la diferencia entre materias es de {difference:.1f} puntos porcentuales.")
        if risk_students:
            interpretations.append(f"Se identifican {len(risk_students)} estudiantes en seguimiento prioritario por una asistencia inferior al 70% o por tener más inasistencias que asistencias.")
        else:
            interpretations.append("No se identifican estudiantes en riesgo alto con los registros disponibles.")
        if classes:
            busiest = max(context_counts, key=context_counts.get)
            interpretations.append(f"El contexto con mayor número de clases registradas es {busiest}, con {context_counts[busiest]} clases.")
        if grades["math"] or grades["science"]:
            averages = []
            if grades["math"]: averages.append(f"Matemáticas {sum(grades['math']) / len(grades['math']):.2f}")
            if grades["science"]: averages.append(f"Ciencias Naturales {sum(grades['science']) / len(grades['science']):.2f}")
            interpretations.append("Promedios de calificación registrados: " + " y ".join(averages) + ".")
        chart_data = {
            "dates": [item[1]["label"] for item in date_items],
            "math": [{"present": item[1]["math_present"], "absent": item[1]["math_absent"]} for item in date_items],
            "science": [{"present": item[1]["science_present"], "absent": item[1]["science_absent"]} for item in date_items],
            "contexts": [{"label": key, "value": value} for key, value in sorted(context_counts.items())],
            "risks": [{"label": item["name"], "value": item["absent"]} for item in risk_students[:8]],
        }
        page_data.update({
            "stats": {"classes": len(classes), "students": len(unique_students), "groups": len({item["group"] for item in classes}), "attendance_rate": overall_rate, "math_rate": math_rate, "science_rate": science_rate, "math_average": round(sum(grades["math"]) / len(grades["math"]), 2) if grades["math"] else None, "science_average": round(sum(grades["science"]) / len(grades["science"]), 2) if grades["science"] else None, "risk_count": len([item for item in student_stats.values() if item["sessions"] and (item["rate"] < 70 or item["absent"] > item["present"])]), "low_performance_count": low_performance_students, "students_with_grades": len(student_grade_averages), "records": len(valid_records), "attendance_coverage": round(records_with_attendance * 100 / len(valid_records), 1) if valid_records else 0, "recent_rate": recent_rate, "trend": trend, "chronic_absence_count": chronic_absence_count, "missing_attendance": missing_attendance, "total_sessions": total_sessions, "total_present": total_present, "total_absent": total_sessions - total_present, "math_sessions": attendance["math"]["sessions"], "science_sessions": attendance["science"]["sessions"]},
            "chart_data": json.dumps(chart_data, ensure_ascii=False), "interpretations": interpretations, "risk_students": risk_students, "probable_exits": probable_exits, "drive_updated": metadata.get("modifiedTime", ""),
        })
        return render_template("seguimiento.html", current_user=session.get("user"), **page_data)
    except Exception:
        app.logger.exception("No se pudo generar seguimiento estadístico")
        page_data["data_error"] = "No se pudo generar el análisis desde Google Drive."
        return render_template("seguimiento.html", current_user=session.get("user"), **page_data)


@app.route("/clases")
@login_required
def mis_clases():
    page_data = {
        "classes": [], "contexts": ["Alta", "Multigrado", "Técnico Laboral", "Comunidad Terapéutica"],
        "subjects": [], "groups": [], "context_filter": "", "subject_filter": "", "group_filter": "", "search": "",
        "data_error": None,
    }
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("clases.html", current_user=session.get("user"), **page_data)
    try:
        buffer, metadata = download_excel_from_drive()
        all_classes = read_planning_rows(buffer)
        context_filter = request.args.get("contexto", "").strip()
        subject_filter = request.args.get("asignatura", "").strip()
        group_filter = request.args.get("grupo", "").strip()
        search = request.args.get("buscar", "").strip()
        contexts = CONTEXT_OPTIONS.copy()
        def compact(value):
            return re.sub(r"[^a-z0-9]", "", normalize_header(value))
        def subject_key(value):
            normalized = compact(value)
            if "matematic" in normalized or normalized in {"mate", "mates"}:
                return "matematicas"
            if "ciencianatural" in normalized or normalized in {"ciencias", "ciencia"}:
                return "cienciasnaturales"
            return normalized
        subjects = sorted({
            "Matemáticas" if subject_key(item["subject"]) == "matematicas" else
            "Ciencias Naturales" if subject_key(item["subject"]) == "cienciasnaturales" else item["subject"]
            for item in all_classes if item["subject"]
        })
        groups = sorted({item["group"] for item in all_classes if item["group"]})
        if context_filter not in contexts: context_filter = ""
        if subject_filter not in subjects: subject_filter = ""
        if group_filter not in groups: group_filter = ""
        needle = compact(search)
        classes = [item for item in all_classes if
            (not context_filter or compact(item["context"]) == compact(context_filter)) and
            (not subject_filter or subject_key(item["subject"]) == subject_key(subject_filter)) and
            (not group_filter or compact(item["group"]) == compact(group_filter)) and
            (not needle or needle in compact(f"{item['subject']} {item['theme']} {item['observations']} {item['group']}"))]
        page_data.update({"classes": classes, "subjects": subjects, "groups": groups, "context_filter": context_filter, "subject_filter": subject_filter, "group_filter": group_filter, "search": search, "drive_updated": metadata.get("modifiedTime", "")})
        return render_template("clases.html", current_user=session.get("user"), **page_data)
    except Exception:
        app.logger.exception("No se pudieron leer las clases")
        page_data["data_error"] = "No se pudo leer el Excel desde Google Drive."
        return render_template("clases.html", current_user=session.get("user"), **page_data)


@app.route("/planeacion")
@login_required
def planeacion():
    page_data = {
        "planning": [], "weekly_groups": [],
        "cleis": ["CLEI 1", "CLEI 2", "CLEI 3", "CLEI 4", "CLEI 5-6"],
        "subjects": [], "weeks": [], "clei_filter": "", "subject_filter": "", "week_filter": "", "search": "", "data_error": None,
    }
    if not DRIVE_ENABLED:
        page_data["data_error"] = "La conexión con Google Drive está pausada."
        return render_template("planeacion.html", current_user=session.get("user"), **page_data)
    try:
        buffer, metadata = download_planning_from_drive()
        planning = read_additional_planning_rows(buffer)
        clei_filter = request.args.get("clei", "").strip()
        subject_filter = request.args.get("materia", "").strip()
        week_filter = request.args.get("semana", "").strip()
        search = request.args.get("buscar", "").strip()
        subjects = sorted({item["subject"] for item in planning})
        weeks = sorted({item.get("week", "") for item in planning if item.get("week")}, key=lambda value: int(re.search(r"\d+", value).group()) if re.search(r"\d+", value) else 999)
        if clei_filter not in page_data["cleis"]: clei_filter = ""
        if subject_filter not in subjects: subject_filter = ""
        if week_filter not in weeks: week_filter = ""
        needle = normalize_header(search)
        visible = [item for item in planning if
            (not clei_filter or item["group"] == clei_filter) and
            (not subject_filter or item["subject"] == subject_filter) and
            (not week_filter or item.get("week") == week_filter) and
            (not needle or needle in normalize_header(f"{item['group']} {item['subject']} {item['theme']}"))]
        weekly = []
        for item in visible:
            # Una semana debe ser una sola fila aunque cada pestaña escriba
            # el rango de fechas con espacios o formato ligeramente distinto.
            week_key = (item.get("week_number", 999), normalize_header(item.get("week", "")))
            if not weekly or weekly[-1]["key"] != week_key:
                weekly.append({"key": week_key, "week": item.get("week", "Semana sin número"), "date_range": item.get("date_label", "Fecha por definir"), "items": [], "cells": {clei: [] for clei in page_data["cleis"]}})
            weekly[-1]["items"].append(item)
            if item.get("group") in weekly[-1]["cells"]:
                weekly[-1]["cells"][item["group"]].append(item)
        page_data.update({"planning": visible, "weekly_groups": weekly, "subjects": subjects, "weeks": weeks, "clei_filter": clei_filter, "subject_filter": subject_filter, "week_filter": week_filter, "search": search, "drive_updated": metadata.get("modifiedTime", "")})
        response = make_response(render_template("planeacion.html", current_user=session.get("user"), **page_data))
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        return response
    except Exception:
        app.logger.exception("No se pudo leer la planeación adicional")
        page_data["data_error"] = "No se pudo leer la hoja adicional de planeación desde Google Drive."
        return render_template("planeacion.html", current_user=session.get("user"), **page_data)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)

