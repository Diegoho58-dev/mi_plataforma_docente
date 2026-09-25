"""Lectura de materias externas desde las planillas operativas de Google Drive.

La salida es de solo lectura y conserva la diferencia entre la hoja/bloque
administrativo y la fecha escrita en el encabezado de cada materia. No lee ni
calcula columnas PROM., ASIST. de resumen ni TOTALES/PROMEDIOS del grupo.
"""

import re
import unicodedata
from datetime import date
from itertools import chain

from openpyxl import load_workbook


OWN_SUBJECT_MARKERS = {
    "matematicas",
    "matematicas fisica",
    "fisica",
    "ciencias naturales",
}

EXTERNAL_SUBJECTS = {
    "espanol": "Español",
    "lengua castellana": "Español",
    "ciencias sociales": "Ciencias Sociales",
    "sociales": "Ciencias Sociales",
    "ingles": "Inglés",
}

CLEI_LABELS = (
    "CLEI 5-6", "CLEI 5", "CLEI 6", "CLEI 3A", "CLEI 3B",
    "CLEI 4", "CLEI 2", "CLEI 1", "MULTIGRADO", "ALF",
)

DATE_RE = re.compile(
    r"(?<!\d)(\d{1,2})\s*(?:/|-|–)\s*(\d{1,2})"
    r"(?:\s*(?:/|-|–)\s*(\d{2,4}))?(?!\d)"
)

SPANISH_DATE_RE = re.compile(
    r"(?<!\d)(\d{1,2})\s*(?:de\s+)?"
    r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|setiembre|"
    r"octubre|noviembre|diciembre)"
    r"(?:\s+(?:de\s+)?(\d{2,4}))?(?!\d)",
    re.IGNORECASE,
)


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split()).strip()


def normalize(value):
    text = clean(value).lower()
    return "".join(
        char for char in unicodedata.normalize("NFD", text)
        if unicodedata.category(char) != "Mn"
    )


def subject_name(value):
    key = normalize(value).replace("/", " ")
    key = re.sub(r"\s+", " ", key).strip()
    for marker, label in EXTERNAL_SUBJECTS.items():
        if marker in key:
            return label
    return ""


def is_own_subject(value):
    key = normalize(value).replace("/", " ")
    return any(marker in key for marker in OWN_SUBJECT_MARKERS)


def parse_date_token(day, month, year=None):
    try:
        day = int(day)
        month = int(month)
        if year is None:
            return None
        year = int(year)
        if year < 100:
            year += 2000
        return date(year, month, day)
    except (TypeError, ValueError):
        return None


def extract_dates(text, default_year=2026):
    """Extrae fechas tolerando /, -, espacios y años de dos dígitos."""
    result = []
    for match in DATE_RE.finditer(clean(text)):
        year = match.group(3) or str(default_year)
        parsed = parse_date_token(match.group(1), match.group(2), year)
        if parsed and parsed not in result:
            result.append(parsed)
    months = {
        "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
        "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
        "septiembre": 9, "setiembre": 9, "octubre": 10,
        "noviembre": 11, "diciembre": 12,
    }
    for match in SPANISH_DATE_RE.finditer(clean(text)):
        year = int(match.group(3) or default_year)
        if year < 100:
            year += 2000
        try:
            parsed = date(year, months[normalize(match.group(2))], int(match.group(1)))
        except (KeyError, TypeError, ValueError):
            parsed = None
        if parsed and parsed not in result:
            result.append(parsed)
    return result


def clei_key(value):
    key = normalize(value).replace(" ", "")
    aliases = {
        "clei1": "CLEI 1", "clei2": "CLEI 2", "clei3a": "CLEI 3A",
        "clei3b": "CLEI 3B", "clei4": "CLEI 4", "clei5": "CLEI 5-6",
        "clei6": "CLEI 5-6", "clei5-6": "CLEI 5-6",
        "1": "CLEI 1", "2": "CLEI 2", "3": "CLEI 3", "3a": "CLEI 3A",
        "3b": "CLEI 3B", "4": "CLEI 4", "5": "CLEI 5-6", "6": "CLEI 5-6",
        "5-6": "CLEI 5-6", "5/6": "CLEI 5-6",
        "multigrado": "MULTIGRADO", "alf": "ALF",
    }
    return aliases.get(key, clean(value).upper())


def student_clei(value, group=""):
    """Obtiene el CLEI operativo y separa CLEI 3 por grupo en Alta."""
    key = clei_key(value)
    group_key = normalize(group)
    if key == "CLEI 3":
        if re.search(r"\bgrupo\s*2\b", group_key):
            return "CLEI 3A"
        if re.search(r"\bgrupo\s*3\b", group_key):
            return "CLEI 3B"
    return key


def context_for(source, group, clei):
    text = normalize(f"{source} {group} {clei}")
    if "multigrado" in text or "mediana" in text:
        return "Multigrado / Mediana"
    if "alf" in text or "alta" in text:
        return "Alta"
    return "CLEI normal"


def sheet_block_label(title, source):
    title_text = clean(title)
    match = re.search(r"semana\s*[-:]?\s*(\d+(?:\s*(?:y|a|-|–)\s*\d+)?)", normalize(title_text))
    if match:
        week_text = re.sub(r"\s+", " ", match.group(1))
        return f"Semana {week_text}"
    return f"Hoja: {title_text or source}"


def _header_row(rows):
    for index, row in enumerate(rows):
        text = " ".join(normalize(cell) for cell in row[:8] if cell)
        if "apellidos y nombres" in text and "identificacion" in text:
            return index
    return None


def _subject_columns(rows, header_index):
    """Encuentra la columna inicial de cada materia en el encabezado visual."""
    found = []
    for row_index in range(max(0, header_index - 12), header_index + 1):
        for column, value in enumerate(rows[row_index]):
            label = subject_name(value)
            if not label or is_own_subject(value):
                continue
            if not any(item[0] == label and item[1] == column for item in found):
                found.append((label, column, row_index))
    # No se deduplican las materias: una misma asignatura puede tener varias
    # clases y cada aparición puede tener una fecha distinta en el encabezado.
    return found


def _teacher_map(rows, header_index):
    mapping = {}
    subject_tokens = sorted(
        set(EXTERNAL_SUBJECTS) | {"matematicas", "ciencias naturales", "fisica"},
        key=len,
        reverse=True,
    )
    subject_pattern = re.compile("|".join(re.escape(item) for item in subject_tokens))
    for row in rows[:header_index]:
        cells = [clean(cell) for cell in row if clean(cell)]
        joined = " ".join(cells)
        if "areas docente" not in normalize(joined):
            continue
        marker_match = re.search(r"areas\s+docente\s+(.+)", normalize(joined))
        if not marker_match:
            continue
        # Usamos el texto original de la fila para conservar tildes en el nombre.
        original = joined
        marker = re.search(r"áreas\s+docente\s+", original, re.IGNORECASE)
        if not marker:
            marker = re.search(r"areas\s+docente\s+", original, re.IGNORECASE)
        if not marker:
            continue
        tail = original[marker.end():].strip()
        first_subject = subject_pattern.search(normalize(tail))
        teacher = tail[:first_subject.start()].strip(" -:") if first_subject else tail
        if not teacher:
            continue
        for token in subject_tokens:
            if token in normalize(tail):
                label = subject_name(token)
                if label:
                    mapping[label] = teacher
    return mapping


def _subject_date_map(text, default_year=2026):
    """Mapea CLEI/MULTIGRADO a la fecha escrita en el encabezado de materia."""
    result = {}
    normalized_text = clean(text)
    label_pattern = "|".join(re.escape(normalize(label)) for label in CLEI_LABELS)
    pattern = re.compile(rf"({label_pattern})\s*:\s*(.*?)(?=(?:{'|'.join(re.escape(normalize(label)) for label in CLEI_LABELS)})\s*:|$)", re.IGNORECASE)
    for match in pattern.finditer(normalize(normalized_text)):
        label = clei_key(match.group(1))
        dates = extract_dates(match.group(2), default_year)
        if dates:
            # ALF queda fuera de la vista actual. Para cada CLEI se conservan
            # todas las fechas explícitas del encabezado: cada una será una
            # columna independiente en la plataforma.
            if label != "ALF":
                result[label] = dates
    # Si el encabezado no asigna fecha a un CLEI concreto, conservamos todas
    # las fechas como referencia, sin inventar la semana académica.
    if not result:
        dates = extract_dates(normalized_text, default_year)
        if dates:
            result["__default__"] = dates
    return result


def _header_text(rows, column, header_index):
    values = []
    for row in rows[max(0, header_index - 12):header_index + 1]:
        if column < len(row) and clean(row[column]):
            values.append(clean(row[column]))
    return " ".join(values)


def _is_data_row(row):
    if len(row) < 5:
        return False
    name = clean(row[1] if len(row) > 1 else "")
    identification = clean(row[2] if len(row) > 2 else "")
    return bool(name and identification and not normalize(name).startswith("total"))


def parse_workbook(buffer, source, default_year=2026):
    """Devuelve registros detallados de materias externas de un XLSX exportado."""
    buffer.seek(0)
    workbook = load_workbook(buffer, data_only=True, read_only=True)
    records = []
    for worksheet in workbook.worksheets:
        # No cargar la hoja completa: las planillas tienen muchas filas y
        # Render puede finalizar el proceso por memoria insuficiente.
        preview = []
        row_iterator = worksheet.iter_rows(values_only=True)
        for _ in range(80):
            try:
                preview.append([clean(value) for value in next(row_iterator)])
            except StopIteration:
                break
        header_index = _header_row(preview)
        if header_index is None:
            continue
        columns = _subject_columns(preview, header_index)
        teachers = _teacher_map(preview, header_index)
        block = sheet_block_label(worksheet.title, source)
        subject_specs = []
        for subject, start_column, _ in columns:
            if is_own_subject(subject):
                continue
            dates_by_clei = _subject_date_map(
                _header_text(preview, start_column, header_index),
                default_year,
            )
            subject_specs.append((subject, start_column, dates_by_clei, teachers.get(subject, "No identificado")))
        # Ya consumimos la vista previa; continuar desde la fila siguiente
        # mantiene constante el uso de memoria y recorre la hoja una sola vez.
        for raw_row in chain(preview[header_index + 1:], row_iterator):
            row = [clean(value) for value in raw_row]
            if not _is_data_row(row):
                continue
            clei = clean(row[3] if len(row) > 3 else "")
            group = clean(row[4] if len(row) > 4 else "")
            key = student_clei(clei, group)
            for subject, start_column, dates_by_clei, teacher in subject_specs:
                if key == "ALF" or (key == "MULTIGRADO" and source == "Alta y CLEI normal"):
                    continue
                class_dates = dates_by_clei.get(key)
                if class_dates is None and len(dates_by_clei) == 1 and "__default__" in dates_by_clei:
                    class_dates = dates_by_clei["__default__"]
                if not class_dates:
                    continue
                attendance = clean(row[start_column] if start_column < len(row) else "")
                grade = clean(row[start_column + 1] if start_column + 1 < len(row) else "")
                # Las columnas PROM./ASIST. y los totales están después de las
                # materias y nunca se recorren como una materia adicional.
                for class_date in class_dates:
                    records.append({
                        "source": source,
                        "sheet": clean(worksheet.title),
                        "block": block,
                        "subject": subject,
                        "teacher": teacher,
                        "context": context_for(source, group, key),
                        "clei": key,
                        "group": group,
                        "student": clean(row[1]),
                        "identification": clean(row[2]),
                        "attendance": attendance,
                        "grade": grade,
                        "observation": "",
                        "class_date": class_date,
                        "class_date_label": class_date.strftime("%d/%m/%Y") if class_date else "No identificada",
                        "week_mismatch": False,
                    })
    workbook.close()
    return records


def parse_values(values, sheet_title, source, default_year=2026):
    """Parsea valores 2D entregados por Google Sheets API, sin descargar XLSX."""
    rows = [[clean(value) for value in row] for row in (values or [])]
    header_index = _header_row(rows)
    if header_index is None:
        return []
    columns = _subject_columns(rows, header_index)
    teachers = _teacher_map(rows, header_index)
    block = sheet_block_label(sheet_title, source)
    subject_specs = []
    for subject, start_column, _ in columns:
        if is_own_subject(subject):
            continue
        dates_by_clei = _subject_date_map(_header_text(rows, start_column, header_index), default_year)
        subject_specs.append((subject, start_column, dates_by_clei, teachers.get(subject, "No identificado")))

    records = []
    for row in rows[header_index + 1:]:
        if not _is_data_row(row):
            continue
        clei = clean(row[3] if len(row) > 3 else "")
        group = clean(row[4] if len(row) > 4 else "")
        key = student_clei(clei, group)
        for subject, start_column, dates_by_clei, teacher in subject_specs:
            if key == "ALF" or (key == "MULTIGRADO" and source == "Alta y CLEI normal"):
                continue
            class_dates = dates_by_clei.get(key)
            if class_dates is None and len(dates_by_clei) == 1 and "__default__" in dates_by_clei:
                class_dates = dates_by_clei["__default__"]
            if not class_dates:
                continue
            attendance = clean(row[start_column] if start_column < len(row) else "")
            grade = clean(row[start_column + 1] if start_column + 1 < len(row) else "")
            for class_date in class_dates:
                records.append({
                    "source": source, "sheet": clean(sheet_title), "block": block,
                    "subject": subject, "teacher": teacher,
                    "context": context_for(source, group, key), "clei": key, "group": group,
                    "student": clean(row[1]), "identification": clean(row[2]),
                    "attendance": attendance, "grade": grade, "observation": "",
                    "class_date": class_date,
                    "class_date_label": class_date.strftime("%d/%m/%Y") if class_date else "No identificada",
                    "week_mismatch": False,
                })
    return records


def consolidate_records(records):
    """Consolida asistencia y nota por hoja, estudiante, CLEI, materia y fecha."""
    consolidated = {}
    for item in records:
        key = (
            item.get("source", ""), item.get("sheet", ""), item.get("subject", ""),
            item.get("clei", ""), item.get("group", ""), item.get("student", ""),
            item.get("identification", ""), item.get("class_date"),
        )
        current = consolidated.get(key)
        if current is None:
            consolidated[key] = dict(item)
            continue
        for field in ("attendance", "grade", "teacher"):
            if not current.get(field) and item.get(field):
                current[field] = item[field]
    return list(consolidated.values())


def mark_week_mismatches(records, cycle_for_date=None):
    """Marca fechas fuera del bloque solo cuando el bloque tiene una semana explícita.

    No mueve registros ni cambia su fecha. Si no se puede determinar la semana
    desde el nombre de la hoja, se deja la marca en False para evitar inferencias.
    """
    if not cycle_for_date:
        return records
    for item in records:
        match = re.search(r"semana\s*(\d+(?:\s*(?:y|a|-|–)\s*\d+)?)", normalize(item["block"]))
        if not match or not item.get("class_date"):
            continue
        cycle = cycle_for_date(item["class_date"])
        expected_weeks = {
            int(value)
            for value in re.findall(r"\d+", match.group(1))
        }
        if cycle and cycle.get("week") not in expected_weeks:
            item["week_mismatch"] = True
    return records


def summarize(records):
    return {
        "records": len(records),
        "teachers": len({item["teacher"] for item in records if item["teacher"] != "No identificado"}),
        "subjects": len({item["subject"] for item in records}),
        "students": len({item["identification"] for item in records if item["identification"]}),
        "dated": sum(bool(item.get("class_date")) for item in records),
        "undated": sum(not item.get("class_date") for item in records),
        "mismatches": sum(item.get("week_mismatch", False) for item in records),
    }


__all__ = [
    "parse_workbook", "parse_values", "consolidate_records", "build_external_matrix", "mark_week_mismatches", "summarize",
    "EXTERNAL_SUBJECTS", "OWN_SUBJECT_MARKERS",
]


def _demo():
    """Solo para pruebas manuales; no se ejecuta al importar."""
    return True


def build_external_matrix(records, source_filter="", subject_filter="", clei_filter="", search=""):
    """Agrupa registros externos por estudiante y fecha para una vista matricial."""
    needle = normalize(search)
    filtered = []
    for item in records:
        if source_filter and item["source"] != source_filter:
            continue
        if subject_filter and item["subject"] != subject_filter:
            continue
        if clei_filter and clei_key(item["clei"]) != clei_key(clei_filter):
            continue
        searchable = normalize(" ".join(str(item.get(key, "")) for key in ("student", "identification", "group", "clei")))
        if needle and needle not in searchable:
            continue
        if not item.get("class_date"):
            continue
        filtered.append(item)

    dates = {}
    students = {}
    for item in filtered:
        date_key = item["class_date"].isoformat()
        dates[date_key] = item["class_date"].strftime("%d/%m/%Y")
        key = (item["student"], item["identification"], item["group"], item["clei"])
        student = students.setdefault(key, {
            "name": item["student"],
            "identification": item["identification"],
            "group": item["group"],
            "clei": item["clei"],
            "dates": {},
        })
        student["dates"].setdefault(date_key, []).append({
            "subject": item["subject"],
            "attendance": item["attendance"],
            "grade": item["grade"],
            "block": item["block"],
            "source": item["source"],
            "week_mismatch": item.get("week_mismatch", False),
        })

    for student in students.values():
        absences = 0
        grades_by_subject = {}
        for cells in student["dates"].values():
            for cell in cells:
                attendance = normalize(cell.get("attendance", ""))
                if attendance in {"no", "no asistio", "ausente", "inasistente"}:
                    absences += 1
                grade_text = clean(cell.get("grade", "")).replace(",", ".")
                try:
                    grades_by_subject.setdefault(cell["subject"], []).append(float(grade_text))
                except (TypeError, ValueError):
                    pass
        student["absences"] = absences
        student["subject_averages"] = {
            subject: round(sum(grades) / len(grades), 2)
            for subject, grades in grades_by_subject.items()
            if grades
        }

    return sorted(students.values(), key=lambda item: item["name"].lower()), sorted(dates.items())
